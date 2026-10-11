"""End-to-end regression for v1.4: WMS stock, blanket, LPN<->pallet dock scan, PDI + double-confirmed sub-inventory
transfer, SO after PDI, invoice + challan, loaded/empty, internal moves, other-plant receipt, empty return + e-way bill,
missed-scan detection. Run:  python -m tests.test_parts_flow   (from server/, uses a temp SQLite db)"""
import os, sys, uuid, tempfile
from datetime import datetime, timedelta
DB = os.getenv("KEEP") or os.path.join(tempfile.gettempdir(), f"at_test_{uuid.uuid4().hex[:6]}.db")
os.environ["PALLET_DB_URL"] = f"sqlite:///{DB}"
os.environ["PALLET_INTEGRATION_MODE"] = "STUB"
os.environ["PALLET_BENCH_SCHEDULER"] = "0"
from fastapi.testclient import TestClient
from app.main import app
from app import models
from app.db import SessionLocal
from app.security import hash_pw

ok_n = 0
def check(cond, msg):
    global ok_n
    if not cond:
        print("FAIL:", msg); sys.exit(1)
    ok_n += 1; print("  ok -", msg)


def setup():
    with SessionLocal() as db:
        db.add_all([models.Plant(code="CHN", name="Chennai", gstin="33AAA", reject_subinv="REJ", wms_stale_hours=4),
                    models.Plant(code="PUN", name="Pune", gstin="27AAA")])
        db.flush()
        db.add(models.Customer(code="HMIL", name="Hyundai", plant_code="CHN"))
        for i in range(1, 8):
            db.add(models.Pallet(pallet_no=f"CHN-P{i:03d}", home_plant="CHN", location_plant="CHN", status="AVAILABLE"))
        for i in range(1, 4):
            db.add(models.Pallet(pallet_no=f"PUN-P{i:03d}", home_plant="PUN", location_plant=None, status="AT_CUSTOMER", customer_code="CHN",
                                 challan_ref="PUN-ST-1", dispatch_date=datetime.utcnow()))
        for uid, plant, roles in (("chn", "CHN", ["FGWH", "QA", "LOG", "SEC_OUT", "SEC_IN", "YARD", "PACK", "PADMIN"]), ("pun", "PUN", ["YARD", "FGWH", "SEC_IN"]),
                                  ("master", None, ["MASTER"])):
            u = models.User(user_id=uid, full_name=uid, plant_code=plant, password_hash=hash_pw("x"), must_change_pw=False,
                            supervisor_allowed=True, supervisor_pin_hash=hash_pw("1234"))
            u.roles = [models.UserRole(role_code=r) for r in roles]
            db.add(u); db.flush()
        db.commit()


def tok(c, u, hht=True):
    r = c.post("/api/v1/auth/login", json={"user_id": u, "password": "x"}); assert r.status_code == 200, r.text
    h = {"Authorization": f"Bearer {r.json()['token']}"}
    if hht:
        h |= {"X-Device": "HHT", "X-Device-Id": "HHT-TEST"}
    return h


def ev(c, h, t, payload):
    r = c.post("/api/v1/sync/push", headers=h, json={"events": [{"event_id": str(uuid.uuid4()), "device_id": "HHT-T", "event_type": t,
                                                                  "payload": payload, "local_ts": datetime.utcnow().isoformat()}]})
    assert r.status_code == 200, r.text
    return r.json()["results"][0]


with TestClient(app) as c:
    setup()
    H, M, P = tok(c, "chn"), tok(c, "master", hht=False), tok(c, "pun")
    W = tok(c, "chn", hht=False)          # same user on the web
    print("0. HHT-only policy")
    r = c.post("/api/v1/sync/push", headers=W, json={"events": []})
    check(r.status_code == 403 and "HHT" in r.text, "web cannot push scans")
    r = c.post("/api/v1/moves", headers=W, json={"scanned": ["CHN-P001"], "to_zone": "PRODUCTION"})
    check(r.status_code == 403, "web cannot move pallets")
    print("1. masters")
    b = c.post("/api/v1/blankets", headers=M, json={"plant_code": "CHN", "customer_code": "HMIL", "part_no": "WS-1", "blanket_no": "BL-9",
                                                    "po_number": "PO-123", "schedule_qty": 100}).json()
    check(b["id"], "blanket saved")
    c.post("/api/v1/subinv-rules", headers=M, json={"plant_code": "CHN", "customer_code": "HMIL", "part_no": "WS-1", "ok_subinv": "HMIL-FG", "reject_subinv": "REJ"})
    print("2. WMS freshness gate")
    r = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "CUSTOMER", "blanket_id": b["id"], "part_qty": 30, "qty": 3})
    check(r.status_code == 400 and "WMS" in r.text, "release blocked until WMS stock uploaded")
    csv = "LPN,PART_NO,QTY,SUBINVENTORY,PALLET_NO\nL1,WS-1,10,FG,CHN-P001\nL2,WS-1,10,FG,CHN-P001\nL3,WS-1,10,FG,CHN-P002\nL4,WS-1,10,FG,\nL5,WS-1,10,FG,\n"
    r = c.post("/api/v1/wms/upload", headers=H, files={"file": ("wms.csv", csv, "text/csv")}).json()
    check(r["new"] == 5, f"WMS upload 5 LPN ({r['batch_id']})")
    pal = c.get("/api/v1/pallets/CHN-P001", headers=H).json()
    check((pal.get("pallet") or pal).get("load_state") == "LOADED", "pallet with 2 LPN shows LOADED")
    print("3. release + reservation")
    r = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "CUSTOMER", "blanket_id": b["id"], "part_qty": 200, "qty": 3})
    check(r.status_code == 400 and "open" in r.text, "blanket schedule qty enforced")
    pk = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "CUSTOMER", "blanket_id": b["id"], "part_qty": 30, "qty": 3}).json()
    no = pk["picklist_no"]
    check(len(pk["reserved_lpns"]) == 3 and pk["so_number"] is None and pk["po_number"] == "PO-123", f"{no}: 3 LPN reserved, PO from blanket, no SO yet")
    print("4. internal moves + zone")
    r = c.post("/api/v1/moves", headers=H, json={"scanned": ["CHN-P001"], "to_zone": "FGWH"}).json()[0]
    check(r["ok"] and r["alerts"] and "YARD -> PRODUCTION" in r["alerts"][0], "Yard->FGWH skips Production: missed move recorded")
    r = c.post("/api/v1/moves", headers=H, json={"scanned": ["CHN-P005"], "to_zone": "YARD", "load": "LOADED"}).json()[0]
    check(not r["ok"] or "already" in r["message"], "Yard refuses loaded pallet")
    print("5. dock: pallet scan fetches all its LPNs; LPN scan fetches its pallet")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no, "scanned": "CHN-P001"})
    check(e["status"] in ("APPLIED", "EXCEPTION") and "L1" in e["result"] and "L2" in e["result"], f"pallet scan -> {e['result']}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no, "scanned": "L3"})
    check(e["status"] == "EXCEPTION" and "MOVE_" in e["result"], "LPN scan fetched CHN-P002 + missed internal move alert")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no, "scanned": "L4", "pallet": "CHN-P003"})
    check(e["status"] == "REJECTED" and "exceeds" in e["result"], "over-pick rejected")
    r = c.post(f"/api/v1/picklists/{no}/confirm", headers=H).json()
    check(r["status"] == "PDI_PENDING", "pick list control -> PDI_PENDING (no SO yet)")
    print("6. PDI with a reject")
    e = ev(c, H, "PDI_MARK", {"picklist_no": no, "scanned": "L5", "result": "OK"})
    check(e["status"] == "REJECTED" and "never dock-scanned" in e["result"], "LPN not dock-scanned caught at PDI")
    for l, res in (("L1", "OK"), ("L2", "REJECT"), ("L3", "OK")):
        c.post(f"/api/v1/picklists/{no}/pdi", headers=H, json={"scanned": l, "result": res, "remarks": "scratch" if res == "REJECT" else None})
    r = c.post(f"/api/v1/picklists/{no}/pdi-complete", headers=H).json()
    check(r["status"] == "OPEN" and r["rejected"] == ["L2"], "reject -> back to OPEN for replacement")
    l2 = c.get("/api/v1/lpns/L2", headers=H).json()
    check(l2["lpn"]["subinventory"] == "REJ" and l2["txns"][0]["status"] == "CONFIRMED" and l2["txns"][0]["ack_at"] and l2["txns"][0]["verified_at"],
          "rejected LPN -> REJ sub-inv, double confirmed")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no, "scanned": "L4", "pallet": "CHN-P001"})
    check(e["status"] in ("APPLIED", "EXCEPTION"), f"replacement L4 linked to CHN-P001: {e['result']}")
    c.post(f"/api/v1/picklists/{no}/confirm", headers=H)
    c.post(f"/api/v1/picklists/{no}/pdi", headers=H, json={"scanned": "L4", "result": "OK"})
    r = c.post(f"/api/v1/picklists/{no}/pdi-complete", headers=H).json()
    check(r["status"] == "READY" and r["so_number"], f"PDI OK -> sub-inv confirmed -> SO {r['so_number']} -> READY")
    l1 = c.get("/api/v1/lpns/L1", headers=H).json()["lpn"]
    check(l1["subinventory"] == "HMIL-FG" and l1["txn_state"] == "CONFIRMED", "OK LPN moved to customer/part sub-inventory")
    bl = [x for x in c.get("/api/v1/blankets", headers=H).json() if x["id"] == b["id"]][0]
    check(bl["released_qty"] == 30 and bl["open_qty"] == 70, "blanket released 30, open 70")
    print("7. invoice + challan + gate")
    r = c.post(f"/api/v1/picklists/{no}/challan", headers=H).json()
    check(r["invoice_no"] and r["challan_no"], f"invoice {r['invoice_no']} fetched, challan {r['challan_no']}")
    c.post(f"/api/v1/picklists/{no}/approve", headers=H, json={"vehicle_no": "TN01AB1234"})
    k = c.get(f"/api/v1/picklists/{no}", headers=H).json()["picklist"]
    r = c.post(f"/api/v1/picklists/{no}/gate-out", headers=H, json={"scanned": k["gcs_no"]}).json()
    check(r["status"] == "DISPATCHED", "dispatched")
    p1 = c.get("/api/v1/pallets/CHN-P001", headers=H).json(); p1 = p1.get("pallet") or p1
    check(p1["status"] == "AT_CUSTOMER" and p1["load_state"] == "LOADED", "pallet AT_CUSTOMER and LOADED")
    pr = c.get(f"/print/challan/{no}?tok=" + H["Authorization"][7:])
    check(pr.status_code == 200 and "L1" in pr.text and r["invoice_no"] in pr.text, "challan print shows LPN lines + invoice")
    print("8. missed scans on return")
    r = c.post("/api/v1/slips", headers=H, json={"customer_code": "HMIL", "mode": "A", "pallets": ["CHN-P001", "CHN-P006"]})
    check(r.status_code == 403 and "customer" in r.text, "Option A slip refused for AIS users (customer HHT only)")
    r = c.post("/api/v1/slips", headers=H, json={"customer_code": "PUN", "mode": "B2", "declared_qty": 3})
    check(r.status_code == 400 and "challan" in r.text, "other-plant return refused on a return slip (challan only)")
    with SessionLocal() as db:
        from app.security import hash_pw as _h
        cu = db.query(models.Customer).filter_by(code="HMIL", plant_code="CHN").first(); cu.api_key_hash = _h("ais_testkey_123456"); cu.api_key_prefix = "ais_test"; db.commit()
    r = c.post("/api/ext/v1/return-slip", headers={"X-API-Key": "ais_testkey_123456"}, json={"pallets": ["CHN-P001", "CHN-P006"], "vehicle_no": "TN09AB0001", "customer_challan_no": "DC-76"}).json()
    check(any("MISSED SCAN (OUT_GATE)" in w for w in r.get("warnings", [])), f"returned pallet never dispatched -> missed OUT_GATE recorded {str(r)[:300]}")
    pk2 = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "CUSTOMER", "blanket_id": b["id"], "part_qty": 10, "qty": 1})
    check(pk2.status_code == 200, "second release ok")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": pk2.json()["picklist_no"], "scanned": "L5", "pallet": "CHN-P002"})
    check(e["status"] == "EXCEPTION" and "RETURN_SLIP" in e["result"] or "YARD_IN" in e["result"], f"pallet still at customer scanned at dock -> auto received: {e['result'][:90]}")
    print("9. other-plant material -> FGWH/Packing -> Yard -> empty return with e-way bill")
    r = c.post("/api/v1/plant-receipts", headers=H, json={"scanned": ["PUN-P001", "PUN-P002"], "zone": "FGWH", "ref": "PUN-ST-1"}).json()
    check(all(x["ok"] for x in r), "PUN pallets received at FGWH (LOADED, IN_WIP)")
    e = ev(c, H, "PALLET_SCAN_YARD", {"scanned": "PUN-P003", "accept_foreign": True})
    check(e["status"] == "REJECTED" and "FGWH" in e["result"], "loaded other-plant pallet refused at Yard")
    r = c.post("/api/v1/moves", headers=H, json={"scanned": ["PUN-P001"], "to_zone": "PACKING"}).json()[0]
    r = c.post("/api/v1/moves", headers=H, json={"scanned": ["PUN-P001", "PUN-P002"], "to_zone": "YARD"}).json()
    check(all(x["ok"] for x in r), "emptied -> Packing -> Yard (PUN-P002 skipped Packing -> recorded)")
    rt = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "EMPTY_RETURN", "to_plant": "PUN", "qty": 2}).json()
    rno = rt["picklist_no"]
    for pn in ("PUN-P001", "PUN-P002"):
        e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": rno, "scanned": pn})
        check(e["status"] in ("APPLIED", "EXCEPTION"), f"{pn} scanned on empty return")
    c.post(f"/api/v1/picklists/{rno}/confirm", headers=H)
    r = c.post(f"/api/v1/picklists/{rno}/challan", headers=H).json()
    check(r["challan_no"] and r["ewaybill_no"] and not r.get("so_number"), f"EBS challan {r['challan_no']} + e-way bill {r['ewaybill_no']}, no SO")
    c.post(f"/api/v1/picklists/{rno}/approve", headers=H, json={"vehicle_no": "TN02X1"})
    k = c.get(f"/api/v1/picklists/{rno}", headers=H).json()["picklist"]
    c.post(f"/api/v1/picklists/{rno}/gate-out", headers=H, json={"scanned": k["challan_no"]})
    g = c.post(f"/api/v1/slips/{k['challan_no']}/gate-in", headers=H)
    check(g.status_code == 400 and "not addressed" in g.text, "challan for PUN refused at CHN IN gate")
    g = c.post("/api/v1/slips/" + f"CHL|{k['challan_no']}|{rno}|2".replace("|", "%7C") + "/gate-in", headers=P)
    check(g.status_code == 200 and g.json()["kind"] == "CHALLAN" and g.json()["from_plant"] == "CHN", f"PUN IN gate accepts CHN challan: {g.text[:90]}")
    e = ev(c, P, "PALLET_SCAN_YARD", {"scanned": "PUN-P001"})
    p = c.get("/api/v1/pallets/PUN-P001", headers=P).json(); p = p.get("pallet") or p
    check(p["status"] == "AVAILABLE" and p["load_state"] == "EMPTY" and p["zone"] == "YARD", "PUN yard receives its empty pallet")
    print("10. WMS re-upload protects in-process LPNs")
    r = c.post("/api/v1/wms/upload", headers=H, files={"file": ("w.csv", "LPN,PART_NO,QTY,SUBINVENTORY\nL5,WS-1,99,XX\nL9,WS-1,10,FG\n", "text/csv")}).json()
    l5 = c.get("/api/v1/lpns/L5", headers=H).json()["lpn"]
    check(r["protected"] == 1 and l5["qty"] == 10 and l5["status"] == "PICKED", "picked LPN protected from WMS overwrite")
    pk3 = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "CUSTOMER", "blanket_id": b["id"], "part_qty": 10, "qty": 1})
    check(pk3.status_code == 200 and pk3.json()["reserved_lpns"] == ["L9"], "L9 reserved")
    r = c.post("/api/v1/wms/upload", headers=H, files={"file": ("w.csv", "LPN,PART_NO,QTY,SUBINVENTORY\nL10,WS-1,10,FG\n", "text/csv")}).json()
    l9 = c.get("/api/v1/lpns/L9", headers=H).json()["lpn"]
    check(r["missing"] >= 1 and l9["status"] == "RESERVED" and l9["wms_missing"], "reserved LPN absent from WMS -> kept reserved + flagged")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": pk3.json()["picklist_no"], "scanned": "L9", "pallet": "CHN-P007"})
    check(e["status"] == "REJECTED" and "WMS" in e["result"], "flagged LPN blocked at dock")
    res = c.get("/api/v1/reservations", headers=H).json()
    check(len(res) >= 1, f"reservation control lists {len(res)} pick list(s)")
    print("10b. customer API + signed QR label")
    with SessionLocal() as db:
        from app.security import hash_pw as _h
        cu = db.query(models.Customer).filter_by(code="HMIL", plant_code="CHN").first(); cu.api_key_hash = _h("ais_testkey_123456"); cu.api_key_prefix = "ais_test"; db.commit()
    K = {"X-API-Key": "ais_testkey_123456"}
    r = c.post("/api/ext/v1/return-slip", headers=K, json={"pallets": ["CHN-P001"], "vehicle_no": "TN 09 AB 1111", "customer_challan_no": "DC-77"}).json()
    check(r["qr_payload"].startswith("AIS1|RS|") and r["label_url"], f"customer slip {r['slip_no']} QR {r['qr_payload']}")
    lab = c.get(r["label_url"]); check(lab.status_code == 200 and "<svg" in lab.text and "TN09AB1111" in lab.text and "PALLET RETURN SLIP" in lab.text, "A4 return slip prints with signed QR")
    check(c.get(r["label_url"][:-3] + "XXX").status_code == 403, "label link cannot be guessed")
    bad = r["qr_payload"].replace("TN09AB1111", "TN09AB9999")
    g = c.post(f"/api/v1/slips/{r['slip_no']}/gate-in", headers=H, json={"qr": bad})
    check(g.status_code == 400 and "check code" in g.text, "altered label rejected at gate")
    g = c.post(f"/api/v1/slips/{r['slip_no']}/gate-in", headers=H, json={"qr": r["qr_payload"], "vehicle_no": "TN09AB1111"})
    check(g.status_code == 200 and g.json()["status"] == "IN_GATE", "genuine label accepted at IN gate")
    print("11. missed-scan register")
    ms = c.get("/api/v1/scan-misses", headers=H).json()
    pts = sorted({m["missed_point"] for m in ms})
    check(len(ms) >= 5, f"{len(ms)} missed scans recorded: {pts}")
    r = c.post(f"/api/v1/scan-misses/{ms[0]['id']}/resolve", headers=H, json={"resolution": "briefed security"}).json()
    check(r["resolved"], "missed scan resolved")
    pp = c.get("/api/v1/reports/plant-position", headers=W).json()
    check(pp["summary"]["own_total"] == 7 and pp["summary"]["at_customers"] >= 1, f"plant position: {pp['summary']}")
    lv = c.get("/api/v1/reports/live", headers=W).json()
    check(len(lv) > 10, f"live feed {len(lv)} rows")
    print("12. user access by central admin only")
    r = c.post("/api/v1/users", headers=W, json={"user_id": "x1", "full_name": "x", "plant_code": "CHN", "roles": ["YARD"], "password": "Abc@1234"})
    check(r.status_code == 403, "plant user cannot create users")
    with SessionLocal() as db:
        u = models.User(user_id="cadm", full_name="central", plant_code=None, password_hash=hash_pw("x"), must_change_pw=False)
        u.roles = [models.UserRole(role_code="CADMIN")]; db.add(u); db.commit()
    A = tok(c, "cadm", hht=False)
    r = c.post("/api/v1/users", headers=A, json={"user_id": "yard1", "full_name": "Yard op", "plant_code": "CHN", "roles": ["YARD"],
                                               "password": "Abc@1234", "perms": ["YARD_SCAN", "MOVE_TO_PRODUCTION"]})
    check(r.status_code == 200, "central admin creates user with exact movement rights")
    with SessionLocal() as db:
        u = db.query(models.User).filter_by(user_id="yard1").first(); u.must_change_pw = False; u.password_hash = hash_pw("x"); db.commit()
    Y = tok(c, "yard1")
    e = ev(c, Y, "PALLET_MOVE", {"scanned": "CHN-P007", "to_zone": "PRODUCTION"})
    check(e["status"] in ("APPLIED", "EXCEPTION"), "allowed route Yard -> Production")
    e = ev(c, Y, "PALLET_MOVE", {"scanned": "CHN-P007", "to_zone": "FGWH"})
    check(e["status"] == "REJECTED" and "MOVE_TO_FGWH" in e["result"], "route not granted -> rejected")
    print("13. Dispatch Planning Bench pull (all customers)")
    from app import bench
    r = c.post("/api/v1/integrations/dispatch-bench", headers=A, json={"bench_enabled": False, "bench_url": "http://bench:8765", "bench_user": "svc",
                                                                   "bench_password": "x", "bench_plant": "CHN", "bench_customer_map": "HVF1=HMIL"})
    check(r.status_code == 200, "bench settings saved by central admin")
    fake = {"/api/dispatches/recent?limit=300": {"dispatches": [{"id": 501, "trip_id": 77, "plant": "HVF1", "vehicle": "TN 22 X 9", "dispatch_status": "Completed", "sent_at": datetime.now().isoformat()},
                                                                   {"id": 502, "trip_id": 78, "customer_code": "MSIL1", "vehicle": "HR55A1", "dispatch_status": "Completed", "sent_at": datetime.now().isoformat()}]},
            "/api/picks?dispatch_id=502": {"picks": []},
            "/api/picks?dispatch_id=501": {"picks": [{"picked_lpn": "BL1", "item_code": "WS-9", "qty": 20, "sub_inv": "HMIL-FG"},
                                                      {"picked_lpn": "BL2", "item_code": "DR-3", "qty": 30, "sub_inv": "HMIL-FG"}]}}
    with SessionLocal() as db:
        res = bench.pull(db, fetch=lambda path: fake[path])
    check(res["created"] == 2 and res["lpns"] == 2, f"bench dispatches 501/502 -> pick lists DB-CHN-501 (2 LPN) and DB-CHN-502 ({res})")
    k2 = c.get("/api/v1/picklists/DB-CHN-502", headers=W).json()["picklist"]
    check(k2["customer_code"] == "MSIL1" and k2["source"] == "DISPATCH_BENCH", "non-Hyundai customer taken from the Bench dispatch")
    with SessionLocal() as db:
        res2 = bench.pull(db, fetch=lambda path: fake[path])
    check(res2["created"] == 0, "second pull updates, no duplicate pick lists")
    check(c.get("/api/v1/integrations/hmil-bench", headers=A).status_code == 200, "old /hmil-bench address still answers")
    pk = c.get("/api/v1/picklists/DB-CHN-501", headers=W).json()["picklist"]
    check(pk["customer_code"] == "HMIL" and pk["part_qty"] == 50 and pk["vehicle_no"] == "TN22X9", "customer map, qty and vehicle from Bench")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": "DB-CHN-501", "scanned": "L10", "pallet": "CHN-P004"})
    check(e["status"] == "REJECTED" and "not on Dispatch Bench trip" in e["result"], "LPN not on the Bench trip refused")
    e1 = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": "DB-CHN-501", "scanned": "BL1", "pallet": "CHN-P004"})
    e2 = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": "DB-CHN-501", "scanned": "BL2", "pallet": "CHN-P004"})
    check(e1["status"] != "REJECTED" and e2["status"] != "REJECTED", f"2 Bench LPN bound to one pallet: {e2['result'][:60]}")
    r = c.post("/api/v1/picklists/DB-CHN-501/confirm", headers=H).json()
    check(r["status"] == "READY", "Bench trip skips AssetTrack PDI/SO -> READY")
    r = c.post("/api/v1/picklists/DB-CHN-501/challan", headers=W).json()
    check(r.get("challan_no"), f"challan {r.get('challan_no')} without EBS invoice wait")
    z = c.get("/api/v1/zones", headers=H).json()
    check("FGWH" in z["zones"], f"zone view: {z['zones']}")
    print("14. Empty pallet return loaded at the Pallet Yard")
    with SessionLocal() as db:
        for pn, zone in (("PUN-P010", "YARD"), ("PUN-P011", "YARD"), ("PUN-P012", "PACKING")):
            db.add(models.Pallet(pallet_no=pn, home_plant="PUN", location_plant="CHN", status="HELD", zone=zone, load_state="EMPTY"))
        db.commit()
    r = c.post("/api/v1/picklists/yard-return", headers=W, json={"to_plant": "PUN"})
    check(r.status_code == 403, "yard loading only from the HHT")
    o = c.get("/api/v1/picklists/yard-return/options", headers=H).json()
    check(any(x["plant"] == "PUN" and x["pallets"] >= 3 for x in o), f"owner plants with pallets held here: {o}")
    k = c.post("/api/v1/picklists/yard-return", headers=H, json={"to_plant": "PUN", "vehicle_no": "tn 09 y 7"}).json()
    yno = k["picklist_no"]
    check(k["load_point"] == "YARD" and k["vehicle_no"] == "TN09Y7" and k["qty"] == 0, f"{yno} started at Yard, open quantity")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": yno, "scanned": "CHN-P005"})
    check(e["status"] == "REJECTED" and "belongs to CHN" in e["result"], "own pallet refused - only the owner plant's pallets")
    for pn in ("PUN-P010", "PUN-P011"):
        e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": yno, "scanned": pn})
        check(e["status"] == "APPLIED", f"{pn} loaded at Yard")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": yno, "scanned": "PUN-P012"})
    check(e["status"] == "EXCEPTION" and "MOVE_PACKING_YARD" in e["result"], f"pallet still shown in Packing -> loaded, missed move recorded: {e['result'][:70]}")
    pull = c.get("/api/v1/sync/pull", headers=H).json()
    check(any(x["picklist_no"] == yno and x["load_point"] == "YARD" for x in pull["picklists"]), "HHT gets the Yard loading list")
    r = c.post(f"/api/v1/picklists/{yno}/confirm", headers=H).json()
    check(r["status"] == "CHALLANED" and r["qty"] == 3 and r["challan_no"] and r["ewaybill_no"], f"Finish loading -> challan {r.get('challan_no')} + e-way bill for 3 pallets")
    p10 = c.get("/api/v1/pallets/PUN-P010", headers=H).json(); p10 = p10.get("pallet") or p10
    check(p10["status"] == "ALLOCATED" and p10["challan_ref"] == r["challan_no"], "pallet carries the challan")
    print("15. sync hardening")
    hb = ev(c, H, "HEARTBEAT", {"pending": 3})
    with SessionLocal() as db:
        n_hb = db.query(models.Event).filter_by(event_type="HEARTBEAT").count()
        dev = db.query(models.Device).filter_by(device_id="HHT-T").first()
    check(hb["status"] == "APPLIED" and n_hb == 0 and dev and dev.pending_reported == 3, "heartbeat updates the device but is not stored as an event")
    eid = str(uuid.uuid4())
    body = {"events": [{"event_id": eid, "device_id": "HHT-T", "event_type": "DAMAGE_MARK", "payload": {"scanned": "CHN-P006"}, "local_ts": datetime.utcnow().isoformat()}]}
    r1 = c.post("/api/v1/sync/push", headers=H, json=body).json()["results"][0]
    r2 = c.post("/api/v1/sync/push", headers=H, json=body).json()["results"][0]
    check(r2.get("duplicate") and r2["status"] == r1["status"], "same event sent twice -> answered from the first result")
    with SessionLocal() as db:
        l = db.get(models.Lpn, "BL1"); l.status = "DISPATCHED"; db.commit()
    d = c.get("/api/v1/sync/pull", headers=H, params={"since": (datetime.utcnow() - timedelta(minutes=5)).isoformat()}).json()
    check(any(x["lpn"] == "BL1" and x["status"] == "DISPATCHED" for x in d["lpns"]), "delta pull reports the LPN that left the live statuses")
    f = c.get("/api/v1/sync/pull", headers=H).json()
    check(not any(x["lpn"] == "BL1" for x in f["lpns"]), "full pull keeps only live LPNs")
    from app.routers.sync import purge_old_events
    with SessionLocal() as db:
        db.query(models.Event).filter_by(event_id=eid).update({"received_ts": datetime.utcnow() - timedelta(days=800)}); db.commit()
        n = purge_old_events(db); db.commit()
    check(n >= 1, f"event purge removed {n} old event(s)")
    from app import config
    check(config.JWT_SECRET != config.LEGACY_SECRET or os.getenv("PALLET_JWT_SECRET"), "server secret is not the public default")
    print("16. AIS pallet QR label (JSON) + auto registration")
    QR = '{"uniquePalletID":"AIS-CHN-ANF-00001-0000001","ownerPlant":"CHN","palletType":"ANF","typeSerial":"00001","globalSerial":"0000001","mfgDate":"2026-08-27"}'
    from app import services
    q = services.parse_pallet_qr(QR)
    check(q and q["pallet_no"] == "AIS-CHN-ANF-00001-0000001" and q["owner_plant"] == "CHN" and q["pallet_type"] == "ANF", "QR JSON parsed")
    e = ev(c, H, "PALLET_SCAN_YARD", {"scanned": QR})
    check(e["status"] == "EXCEPTION" and "registered" in e["result"], f"first scan registers the pallet: {e['result'][:70]}")
    r = c.get("/api/v1/pallets/AIS-CHN-ANF-00001-0000001", headers=H).json(); pal = r.get("pallet") or r
    check(pal["status"] == "AVAILABLE" and pal["pallet_type"] == "ANF" and pal["zone"] == "YARD" and pal["home_plant"] == "CHN", "new pallet AVAILABLE in Yard with type from QR")
    e = ev(c, H, "PALLET_SCAN_YARD", {"scanned": QR})
    check(e["status"] in ("APPLIED", "EXCEPTION") and "registered" not in e["result"], "second scan is a normal yard scan")
    r = c.get("/api/v1/pallets/" + QR.replace("{", "%7B").replace("}", "%7D").replace('"', "%22"), headers=H)
    check(r.status_code == 200, "lookup by scanning the QR works")
    QR2 = QR.replace("CHN", "PUN").replace("0000001", "0000002")
    e = ev(c, H, "PALLET_SCAN_YARD", {"scanned": QR2, "accept_foreign": True})
    r = c.get("/api/v1/pallets/AIS-PUN-ANF-00001-0000002", headers=H).json(); pal = r.get("pallet") or r
    check(pal["home_plant"] == "PUN" and pal["status"] == "HELD" and pal["location_plant"] == "CHN", "other plant's new pallet registered as HELD here")
    QR3 = QR.replace('"CHN"', '"ZZZ"').replace("AIS-CHN", "AIS-ZZZ")
    e = ev(c, H, "PALLET_SCAN_YARD", {"scanned": QR3})
    check(e["status"] == "EXCEPTION" and "not in Plant master" in e["result"], "unknown owner plant -> quarantined with reason")
    e = ev(c, H, "PALLET_MOVE", {"scanned": QR, "to_zone": "PRODUCTION"})
    check(e["status"] in ("APPLIED", "EXCEPTION"), "internal move by scanning the QR")
    print("17. GCS inbox -> loading sheet -> finish loading (challan) -> GCS QR -> OUT gate")
    from app import gcs as gcssvc
    with SessionLocal() as db:
        gcssvc.save_settings(db, {"gcs_enabled": False, "gcs_plant": "CHN", "gcs_host": "ftp.test"}); db.commit()
        for i in (31, 32, 33):
            db.add(models.Pallet(pallet_no=f"CHN-P0{i}", home_plant="CHN", location_plant="CHN", status="AVAILABLE", zone="FGWH", load_state="LOADED"))
        db.commit()
    files = [("GCS-CHN-7781.json", b'{"GCS No":"GCS-CHN-7781","Vehicle":"TN 09 AB 1234","Customer Code":"HMIL","Invoice No":"INV/26/0091","Invoice Date":"2026-10-10","Qty":60,"Item":"WS-1","SO No":"SO-9981"}'),
             ("batch.csv", b"gcs_no,vehicle,customer,invoice no,PDI By,PDI Login,Shift Supervisor,Supervisor Login\nGCS-CHN-7782,KA01ZZ9999,MSIL1,INV/26/0092,Ravi Kumar,ravi.k,Suresh Babu,suresh.b\n"),
             ("gatepass_7783.txt", b"GATE PASS\nGCS No : GCS-CHN-7783\nVehicle : MH 12 CD 4455\nConsignee code : HMIL\nInvoice No : INV/26/0093\n")]
    with SessionLocal() as db:
        res = gcssvc.pull(db, fetch=lambda: files)
    check(res["created"] == 3, f"3 GCS files (json / csv / text) -> 3 loading sheets: {res}")
    k = c.get("/api/v1/picklists/GCS-CHN-GCSCHN7781", headers=W).json()["picklist"]
    check(k["vehicle_no"] == "TN09AB1234" and k["invoice_no"] == "INV/26/0091" and k["customer_code"] == "HMIL" and k["source"] == "GCS", "loading sheet carries vehicle, invoice, customer from the GCS file")
    k2 = c.get("/api/v1/picklists/GCS-CHN-GCSCHN7782", headers=W).json()["picklist"]
    check(k2["pdi_sign"] == "Ravi Kumar (ravi.k)" and k2["supervisor_sign"] == "Suresh Babu (suresh.b)", "PDI and shift supervisor (Bench logins) taken from the CSV")
    pr2 = c.get("/api/v1/picklists/GCS-CHN-GCSCHN7782/gcs.pdf", headers=W)
    t2 = "\n".join(pg.extract_text() for pg in __import__('pypdf').PdfReader(__import__('io').BytesIO(pr2.content)).pages)
    check("Ravi Kumar" in t2 and "Suresh Babu" in t2 and "LOGISTICS" in t2 and "SECURITY" in t2 and "DRIVER" not in t2, "PDF shows PDI and Shift Supervisor sign boxes with the Bench names")
    with SessionLocal() as db:
        res = gcssvc.pull(db, fetch=lambda: files)
    check(res["created"] == 0 and res["updated"] == 3, "same files again -> updated, no duplicates")
    r = c.post("/api/v1/picklists/gcs-open", headers=H, json={"gcs_no": "GCS-CHN-7781"}).json()
    check(r["picklist_no"] == "GCS-CHN-GCSCHN7781", "HHT opens the loading sheet by GCS number")
    r = c.post("/api/v1/picklists/gcs-open", headers=H, json={"gcs_no": "GCS-CHN-9999"})
    check(r.status_code == 404, "unknown GCS asks for manual details")
    r = c.post("/api/v1/picklists/gcs-open", headers=H, json={"gcs_no": "GCS-CHN-9999", "vehicle_no": "tn 1 x 1", "customer_code": "HMIL"}).json()
    check(r["gcs_no"] == "GCS-CHN-9999" and r["vehicle_no"] == "TN1X1", "manual GCS opened from the HHT")
    no = "GCS-CHN-GCSCHN7781"
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no, "scanned": "CHN-P031"})
    check(e["status"] in ("APPLIED", "EXCEPTION"), f"pallet tag loaded on the GCS: {e['result'][:60]}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no, "scanned": "PC-000123", "pallet": "CHN-P032"})
    check(e["status"] in ("APPLIED", "EXCEPTION") and "PC-000123" in e["result"], f"unknown part card + pallet loaded: {e['result'][:60]}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no, "scanned": "CHN-P033"})
    r = c.post(f"/api/v1/picklists/{no}/confirm", headers=H).json()
    check(r["status"] == "APPROVED" and r["challan_no"] and r["gcs_qr"].startswith("AIS1|GCS|GCS-CHN-7781|CHN|HMIL|TN09AB1234|3|"), f"finish loading -> challan {r.get('challan_no')} + signed GCS QR, ready for the gate")
    qr = r["gcs_qr"]
    pr = c.get(f"/print/gcs/{no}?tok=" + W["Authorization"].split()[1])
    check(pr.status_code == 200 and "GCS-CHN-7781" in pr.text and "TN09AB1234" in pr.text, "GCS print page with QR")
    bad = qr[:-8] + "DEADBEEF"
    r = c.post("/api/v1/gate-out/scan", headers=H, json={"scanned": bad})
    check(r.status_code == 400 and "HOLD" in r.text, "forged GCS QR refused at the gate")
    r = c.post("/api/v1/gate-out/scan", headers=H, json={"scanned": qr}).json()
    check(r["status"] == "DISPATCHED", "OUT gate: GCS QR -> dispatched")
    p31 = c.get("/api/v1/pallets/CHN-P031", headers=H).json(); p31 = p31.get("pallet") or p31
    check(p31["status"] == "AT_CUSTOMER" and p31["customer_code"] == "HMIL" and p31["challan_ref"] == r["challan_no"], "pallets AT_CUSTOMER with the challan")
    with SessionLocal() as db:
        l = db.get(models.Lpn, "PC-000123")
    check(l and l.status == "DISPATCHED", "part card marked dispatched")
    print("18. GCS PDF: QR + challan + security sign; auto print setting")
    with SessionLocal() as db:
        pl = db.get(models.Plant, "CHN"); pl.gcs_print_mode = "RAW9100"; pl.gcs_printer = "127.0.0.1:1"; db.commit()
    r = c.get(f"/api/v1/picklists/{no}/gcs.pdf", headers=W)
    check(r.status_code == 200 and r.content[:4] == b"%PDF", "GCS PDF downloads")
    from pypdf import PdfReader
    import io as _io
    text = "\n".join(pg.extract_text() for pg in PdfReader(_io.BytesIO(r.content)).pages)
    check("GCS-CHN-7781" in text and "CHN/00003" in text and "chn" in text and "SECURITY" in text, "PDF carries GCS no, challan no and the security sign with login ID")
    pr = c.post(f"/api/v1/picklists/{no}/gcs-print", headers=W).json()
    check(pr["pdf"].startswith("GCS_") and "queued" in pr["print"], f"reprint builds the PDF and queues it on the plant printer: {pr}")
    import time as _t; _t.sleep(1.5)
    with SessionLocal() as db:
        from app.integration import IntegrationLog
        lg = db.query(IntegrationLog).filter_by(system="PRINT").order_by(IntegrationLog.id.desc()).first()
    check(lg is not None and not lg.ok, "print attempt logged (test printer unreachable -> failure recorded, not hidden)")
    print("19. One GCS with several invoices - invoice / item / qty at line level; short pallet numbers on the sheet")
    multi = (b"GCS No,Vehicle No,Vehicle Type,Customer Code,Invoice No,Invoice Date,SO No,Item Code,Item Description,Qty,Pallets\n"
             b"GCS-CHN-9001,TN09XY4321,32 FT MXL,MSIL1,INV/26/0201,10-10-2026,SO-501,WS-1,Windshield front,40,2\n"
             b"GCS-CHN-9001,TN09XY4321,32 FT MXL,MSIL1,INV/26/0201,10-10-2026,SO-501,BL-9,Backlite,20,1\n"
             b"GCS-CHN-9001,TN09XY4321,32 FT MXL,MSIL1,INV/26/0202,10-10-2026,SO-502,WS-1,Windshield front,10,1\n"
             b"GCS-CHN-9001,TN09XY4321,32 FT MXL,MSIL1,INV/26/0203,10-10-2026,SO-503,DG-4,Door glass,30,1\n")
    with SessionLocal() as db:
        res = gcssvc.pull(db, fetch=lambda: [("multi.csv", multi)])
    check(res["created"] == 1, f"4 CSV rows of one GCS -> ONE loading sheet: {res}")
    no9 = "GCS-CHN-GCSCHN9001"
    d = c.get(f"/api/v1/picklists/{no9}", headers=W).json()
    it = d["items"]
    check(len(it) == 4 and [x["invoice_no"] for x in it] == ["INV/26/0201", "INV/26/0201", "INV/26/0202", "INV/26/0203"], "4 invoice lines kept in order")
    check([x["part_no"] for x in it] == ["WS-1", "BL-9", "WS-1", "DG-4"] and [x["qty"] for x in it] == [40, 20, 10, 30], "item and qty at line level")
    check(d["picklist"]["part_qty"] == 100 and d["picklist"]["qty"] == 5 and d["picklist"]["invoice_no"] == "INV/26/0201 +2", f"header = totals + invoice summary: {d['picklist']['invoice_no']}")
    with SessionLocal() as db:                          # re-sent file with a correction replaces the lines (still OPEN)
        gcssvc.pull(db, fetch=lambda: [("multi2.csv", multi.replace(b",30,1", b",35,1"))])
    it = c.get(f"/api/v1/picklists/{no9}", headers=W).json()["items"]
    check(len(it) == 4 and it[3]["qty"] == 35, "re-sent GCS file refreshes the lines, no duplicates")
    with SessionLocal() as db:                          # loaded qty fills same-item lines in order
        pk9 = db.get(models.PickList, no9)
        db.add(models.PickListLpn(picklist_no=no9, lpn_no="PC-X1", pallet_no="AIS-CHN-ANF-00007-0000123", part_no="WS-1", qty=45, pdi_result="OK"))
        db.add(models.PickListLine(picklist_no=no9, pallet_no="AIS-CHN-ANF-00007-0000123")); db.commit()
        il = gcssvc.items_loaded(db, no9)
        check([x["loaded"] for x in il] == [40, 0, 5, 0], f"loaded WS-1 45 -> line 1 full 40, line 3 gets 5: {[x['loaded'] for x in il]}")
        check(gcssvc.short_pallet("AIS-CHN-ANF-00007-0000123") == "ANF-00007" and gcssvc.short_pallet("CHN-P001") == "CHN-P001", "short pallet = type + serial; other numbers unchanged")
        from app import gcsprint
        pdf = gcsprint.build_pdf(db, pk9, None, "AIS1|GCS|test")
    rd = PdfReader(_io.BytesIO(pdf)); text = rd.pages[0].extract_text()
    check(len(rd.pages) == 1 and all(x in text for x in ("INV/26/0201", "INV/26/0202", "INV/26/0203", "BL-9", "Inv No.", "Cust Part No", "Line Amount", "Finished Goods Materials Gate Pass", "Vehicle Type", "32 FT MXL", "ANF-00007")) and "AIS-CHN-ANF-00007" not in text,
          "PDF: one A4 page with every invoice line and the short pallet number")
    h = c.get(f"/print/gcs/{no9}?tok={W['Authorization'].split()[1]}").text
    check("INV/26/0203" in h and "Cust Part No" in h and "ANF-00007" in h, "web GCS print lists the invoice lines")
    with SessionLocal() as db:                          # long load -> several pages, "X of Y", signs on the last page
        for k in range(70):
            db.add(models.PickListItem(picklist_no=no9, line_no=10 + k, invoice_no=f"INV/9/{k}", part_no="WS-1", qty=5, amount=1000.0))
        for k in range(120):
            db.add(models.PickListLine(picklist_no=no9, pallet_no=f"AIS-CHN-ANF-{k + 100:05d}-{k + 5000:07d}"))
        db.commit()
        rd = PdfReader(_io.BytesIO(gcsprint.build_pdf(db, db.get(models.PickList, no9), None, "AIS1|GCS|test")))
    last = rd.pages[-1].extract_text()
    check(len(rd.pages) >= 3 and f"{len(rd.pages)} of {len(rd.pages)}" in last and "SECURITY" in last and "ANF-00219" in "".join(p.extract_text() for p in rd.pages),
          f"74 lines + 121 pallets -> {len(rd.pages)} pages, every line / pallet printed, signs on the last page")
    with SessionLocal() as db:                          # returnable pallet challan PDF (Rule 55): 3 copies, HSN / value from plant
        pl = db.get(models.Plant, "CHN"); pl.pallet_hsn = "7326"; db.commit()
        types_on_list = sorted({(db.get(models.Pallet, l.pallet_no).pallet_type or "PALLET").upper() for l in db.query(models.PickListLine).filter_by(picklist_no=no)})
    r = c.post("/api/v1/pallet-types", headers=M, json={"code": types_on_list[0].lower(), "description": "Steel returnable rack", "value": 4500.5,
                                                                "length_mm": 1200, "width_mm": 1000, "height_mm": 1100, "weight_kg": 85, "colour": "Blue"})
    check(r.status_code == 200 and r.json()["code"] == types_on_list[0] and r.json()["value"] == 4500.5, f"pallet type master saved (code upper-cased): {types_on_list[0]}")
    check(c.post("/api/v1/pallet-types", headers=M, json={"code": "X1", "value": -5}).status_code == 400, "negative value refused")
    pt = c.get("/api/v1/pallet-types", headers=W).json()
    check(any(t["code"] == types_on_list[0] for t in pt["types"]) and "missing" in pt, f"pallet types listed; types without master: {pt['missing']}")
    r = c.get(f"/api/v1/picklists/{no}/challan.pdf", headers=W)
    rd = PdfReader(_io.BytesIO(r.content)); ct = "\n".join(pg.extract_text() for pg in rd.pages)
    check(r.status_code == 200 and all(x in ct for x in ("ORIGINAL FOR CONSIGNEE", "DUPLICATE FOR TRANSPORTER", "TRIPLICATE FOR CONSIGNOR", "CHN/00003", "7326", "Rule 55", "Steel returnable rack", "1200x1000x1100 mm", "85 kg", "Blue", "4,500.50", "Total Pallet Weight")),
          f"challan PDF: 3 copies, HSN, description + value per pallet from the Pallet Type master ({len(rd.pages)} pages)")
    check(c.get("/api/v1/picklists/NOPE/challan.pdf", headers=W).status_code == 404, "challan PDF refused without a challan")
    check(gcsprint.amount_in_words(223580.88) == "Two Lakh Twenty Three Thousand Five Hundred Eighty and paise Eighty Eight only", "amount in words (Indian numbering)")
    print("20. Real ERP GCS files (one row per case / LPN): merge to invoice lines, expected part cards, loading check")
    fx = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
    real = [(n, open(os.path.join(fx, n), "rb").read()) for n in sorted(os.listdir(fx)) if n.endswith(".csv")]
    with SessionLocal() as db:
        db.add(models.Customer(code="WEBASTO", name="Webasto Roofsystems India Private Limited", plant_code="CHN"))
        for i in range(41, 47):
            db.add(models.Pallet(pallet_no=f"CHN-P0{i}", home_plant="CHN", location_plant="CHN", status="AVAILABLE", zone="FGWH"))
        # pallet QRs: type + serial match the ERP reference (F.2948 -> AIS-CHN-F-02948-..., NF7944 -> AIS-CHN-NF-07944-...)
        for pno in ("AIS-CHN-F-02948-0010001", "AIS-CHN-F-03655-0010002", "AIS-CHN-NF-07944-0010003", "AIS-PUN-F-02948-0020001"):
            db.add(models.Pallet(pallet_no=pno, home_plant=pno.split("-")[1], location_plant="CHN", status="AVAILABLE", zone="FGWH",
                                 pallet_type=pno.split("-")[2]))
        db.commit()
        res = gcssvc.pull(db, fetch=lambda: real)
        g1 = db.query(models.GcsLpn).filter_by(lpn_no="P10266528972").first()
        g2 = db.query(models.GcsLpn).filter_by(lpn_no="P10266528888").first()
        l1 = db.get(models.Lpn, "P10266528972")
    check(res["created"] == 2, f"2 real GCS files -> 2 loading sheets: {res}")
    check(g1.pallet_no == "AIS-CHN-F-02948-0010001" and l1.pallet_no == g1.pallet_no and g2.pallet_no is None,
          "F.2948 -> own plant's pallet AIS-CHN-F-02948 (not PUN's); part card linked to it; unknown pallet ref left open")
    d = c.get("/api/v1/picklists/GCS-CHN-FGCHN262712814", headers=W).json(); k = d["picklist"]; it = d["items"]
    check(len(it) == 1 and it[0]["cases"] == 6 and it[0]["qty_per_case"] == 120 and it[0]["qty"] == 720 and abs(it[0]["amount"] - 1105304.16) < 0.01,
          f"6 case rows -> 1 invoice line: 6 cases x 120 = 720, amount {it[0]['amount'] if it else '-'}")
    check(k["qty"] == 6 and k["vehicle_no"] == "TN87B1498" and k["vehicle_type"] == "LCV23" and k["customer_code"] == "WEBASTO" and k["customer_location"] == "KANCHEEPURAM"
          and k["gp_date"].startswith("2026-10-09T14:42:22") and k["transporter_name"] == "SRI RAGAVENDRA TRANSPORT" and it[0]["cust_part"] == "5711605A",
          "header: 6 pallets expected, vehicle, type, customer matched by name, location, entry time, transporter, customer part")
    k2 = c.get("/api/v1/picklists/GCS-CHN-FGCHN262712823", headers=W).json()["picklist"]
    check(k2["qty"] == 2 and k2["customer_name"].startswith("RENAULT NISSAN") and k2["invoice_no"] == "604118430 +1", "12823: 2 invoices, 2 pallets, customer name kept")
    no20 = "GCS-CHN-FGCHN262712814"
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no20, "scanned": "AIS-CHN-F-02948-0010001"})
    check(e["status"] == "APPLIED" and "P10266528972" in e["result"], f"pallet QR scanned -> its part card from the GCS picked automatically: {e['result'][:70]}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no20, "scanned": "F.3655.P10266528883"})
    check(e["status"] == "APPLIED" and "AIS-CHN-F-03655" in e["result"] and "x120" in e["result"], f"part card label alone -> pallet already known: {e['result'][:70]}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no20, "scanned": "P10266512005", "pallet": "CHN-P042"})
    check(e["status"] == "REJECTED" and "wrong vehicle" in e["result"], f"part card of GCS 12823 refused on 12814: {e['result'][:70]}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no20, "scanned": "AIS-CHN-NF-07944-0010003"})
    check(e["status"] == "REJECTED" and "wrong vehicle" in e["result"] and "NF-07944" in e["result"], f"pallet of GCS 12823 refused on 12814: {e['result'][:70]}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": no20, "scanned": "P99999999", "pallet": "CHN-P043"})
    msg = (e.get("result") or "") + str(e.get("alerts") or "")
    check(e["status"] == "REJECTED" and "NOT in GCS" in msg and "do PDI" in msg, f"part card not in the GCS file -> stopped, do PDI: {msg[:120]}")
    r = c.post(f"/api/v1/picklists/{no20}/confirm", headers=H).json()
    gc = r.get("gcs_check") or {}
    check(r["status"] == "APPROVED" and len(gc.get("missing_lpns", [])) == 4 and gc.get("extra_lpns") == [] and gc.get("expected_pallets") == 6
          and gc.get("missing_pallets") == [] and gc.get("extra_pallets") == [],
          f"finish loading reports the gap: 4 part cards not scanned, nothing extra (non-PDI card was stopped): {gc}")
    check(gc.get("invoice_qty") == 720 and gc.get("loaded_qty") == 240 and any(d["diff"] == -480 for d in gc.get("qty_diff", [])),
          f"finish loading compares quantity: invoice 720, loaded 240 (short 480): {gc.get('qty_diff')}")
    rd = PdfReader(_io.BytesIO(c.get(f"/api/v1/picklists/{no20}/gcs.pdf", headers=W).content)); gt = rd.pages[0].extract_text()
    check(all(x in gt for x in ("FGCHN2627/12814", "WEBASTO ROOFSYSTEMS", "LCV23", "5711605A", "720.00", "1,105,304.16")), "gate pass PDF from the real file")
    hdr, *rows = real[0][1].decode().splitlines()
    cols = hdr.split(","); ci = cols.index("CANCELLED_DATE")
    import csv as _csv
    rr = list(_csv.reader(rows)); rr = [x for x in rr if x]
    for x in rr:
        x[ci] = "10-OCT-2026"
    buf = _io.StringIO(); w = _csv.writer(buf); w.writerow(cols); w.writerows(rr)
    with SessionLocal() as db:
        res = gcssvc.pull(db, fetch=lambda: [("cancel.csv", buf.getvalue().encode())])
    k2 = c.get("/api/v1/picklists/GCS-CHN-FGCHN262712823", headers=W).json()["picklist"]
    check(res.get("cancelled") == 1 and k2["status"] == "CANCELLED", f"GCS cancelled in the ERP (CANCELLED_DATE) -> loading sheet cancelled: {res}")
    print("21. Masters: auto-created from GCS, Excel template / upload / delete, table export, per-plant FTP folder")
    cu = {x["code"]: x for x in c.get("/api/v1/customers?plant=CHN", headers=M).json()}
    trs = {x["code"]: x for x in c.get("/api/v1/transporters", headers=M).json()}
    check("NISSAN" in cu and cu["NISSAN"]["name"].startswith("RENAULT NISSAN") and cu["NISSAN"]["city"] == "ORAGADAM" and cu["NISSAN"]["ctype"] == "OEM",
          f"customer NISSAN created from the GCS (code = short name in REMARKS, name / location / OEM from the file)")
    check("ARC" in trs and "TN87D2115" in (trs["ARC"]["vehicles"] or "") and "SRIRAGAVENDR" in trs and "CHN" in (trs["ARC"]["plant_codes"] or ""),
          f"transporters created from the GCS with vehicle and plant: {sorted(trs)}")
    k2 = c.get("/api/v1/picklists/GCS-CHN-FGCHN262712823", headers=W).json()["picklist"]
    check(k2["customer_code"] == "NISSAN" and k2["transporter_code"] == "ARC", "loading sheet linked to the new customer / transporter codes")
    from openpyxl import load_workbook, Workbook
    for kind in ("plants", "customers", "transporters", "pallet-types", "pallets"):
        r = c.get(f"/api/v1/masters/{kind}/template.xlsx", headers=M)
        wb = load_workbook(_io.BytesIO(r.content))
        check(r.status_code == 200 and wb.worksheets[0].max_row == 3 and "How to fill" in wb.sheetnames, f"{kind}: template with headings + 2 sample rows + help sheet")
    r = c.get("/api/v1/masters/customers/template.xlsx", headers=M)
    res = c.post("/api/v1/masters/customers/upload", headers=M, files={"file": ("t.xlsx", r.content)}).json()
    check(res["inserted"] == 0 and res["skipped"] == 2, f"untouched template uploads nothing (samples skipped): {res}")
    wb = Workbook(); ws = wb.active
    ws.append(["Code *", "Name (as on GCS) *", "Plant *", "Type (OEM/OES/AFM)", "City", "GSTIN", "Hold Days", "Chargeable (Y/N)", "Active (Y/N)"])
    ws.append(["tvsm", "TVS MOTOR COMPANY LTD", "CHN", "OEM", "Hosur", "33AAACT0000A1Z1", 30, "Y", "Y"])
    ws.append(["NISSAN", "RENAULT NISSAN AUTOMOTIVE INDIA PVT. LTD.", "CHN", "OEM", "Oragadam", "33AAACR0000A1Z9", 45, "N", "Y"])
    ws.append(["BAD", "", "CHN", "", "", "", "x", "", ""])
    ws.append(["XYZ", "Unknown plant cust", "ZZZ", "", "", "", "", "", ""])
    b = _io.BytesIO(); wb.save(b)
    res = c.post("/api/v1/masters/customers/upload", headers=M, files={"file": ("c.xlsx", b.getvalue())}).json()
    cu = {x["code"]: x for x in c.get("/api/v1/customers?plant=CHN", headers=M).json()}
    check(res["inserted"] == 1 and res["updated"] == 1 and len(res["errors"]) == 2 and "TVSM" in cu and cu["NISSAN"]["gstin"] == "33AAACR0000A1Z9"
          and cu["NISSAN"]["holding_limit_days"] == 45 and cu["NISSAN"]["chargeable"] is False, f"upload: 1 added, 1 updated, 2 row errors listed: {res['errors']}")
    r = c.get("/api/v1/masters/customers/data.xlsx", headers=M)
    check(r.status_code == 200 and "TVSM" in str([x for x in load_workbook(_io.BytesIO(r.content)).worksheets[0].iter_rows(values_only=True)]), "current customers downloaded in the template layout")
    d = c.delete(f"/api/v1/masters/customers/{cu['TVSM']['id']}", headers=M).json()
    d2 = c.delete(f"/api/v1/masters/customers/{cu['NISSAN']['id']}", headers=M).json()
    cu = {x["code"]: x for x in c.get("/api/v1/customers?plant=CHN", headers=M).json()}
    check(d["result"] == "deleted" and "TVSM" not in cu and d2["result"] == "deactivated" and cu["NISSAN"]["active"] is False,
          f"delete: unused customer removed; customer with gate passes set inactive ({d2['message']})")
    dp = c.delete("/api/v1/masters/pallets/CHN-P001", headers=M).json()
    check(dp.get("result") == "deactivated" or "cannot be removed now" in str(dp), f"pallet with history -> DISCONTINUED or refused while out (never removed): {dp}")
    check(c.delete("/api/v1/masters/customers/1", headers=H).status_code == 403, "delete needs the master-data right")
    body = {"title": "Pallets test", "columns": ["Pallet", "Dispatched", "Qty"], "rows": [["CHN-P001", "2026-10-09T14:42:22", "120"], ["CHN-P002", "", "7"]]}
    rx = c.post("/api/v1/export/xlsx", headers=W, json=body); rp = c.post("/api/v1/export/pdf", headers=W, json=body)
    ws = load_workbook(_io.BytesIO(rx.content)).worksheets[0]
    check(rx.status_code == 200 and ws.cell(row=5, column=2).number_format.startswith("dd-mmm-yy") and ws.cell(row=5, column=3).value == 120
          and rp.content[:4] == b"%PDF" and "09-Oct-26" in PdfReader(_io.BytesIO(rp.content)).pages[0].extract_text(), "table export: Excel (dates dd-mmm-yy, numbers as numbers) and PDF")
    r = c.post("/api/v1/integrations/gcs-inbox/plant", headers=A, json={"plant": "PUN", "gcs_enabled": True, "gcs_folder": "/GCS/PUN", "gcs_user": "pun_ftp", "gcs_password": "s3cret"})
    check(r.status_code == 200, "per-plant FTP folder saved")
    g = c.get("/api/v1/integrations/gcs-inbox", headers=A).json()
    pun = [x for x in g["plants"] if x["code"] == "PUN"][0]
    check(pun["gcs_folder"] == "/GCS/PUN" and pun["gcs_enabled"] and pun["has_password"] and "s3cret" not in str(g), "plant folder listed, password never shown")
    with SessionLocal() as db:
        cfg = gcssvc.plant_cfg(db, db.get(models.Plant, "PUN"))
        check(cfg["gcs_folder"] == "/GCS/PUN" and cfg["gcs_user"] == "pun_ftp" and cfg["gcs_host"] == "ftp.test", "plant uses its own folder / login and the shared server")
        pun_csv = real[1][1].replace(b"FGCHN2627/12814", b"FGPUN2627/00001")
        res = gcssvc.pull(db, fetch=lambda: [("pun.csv", pun_csv)])
        k = db.query(models.PickList).filter_by(gcs_no="FGPUN2627/00001").first()
    check(res["plants"]["PUN"]["created"] == 1 and k and k.plant_code == "PUN", f"automatic read goes plant by plant: the PUN folder's file opens a PUN loading sheet: {res['plants']}")
    print("22. Customer HHT (Option A at the customer): pairing QR, device key, AIS network / internet route, revoke")
    import base64 as _b64, json as _json
    with SessionLocal() as db:
        hm = db.query(models.Customer).filter_by(code="HMIL", plant_code="CHN").first(); hm_id = hm.id
        held = [x.pallet_no for x in db.query(models.Pallet).filter_by(customer_code="HMIL", status="AT_CUSTOMER").all()]
    check(c.get("/api/ext/v1/ping").json()["ok"], "customer API answers /ping without a key (route test)")
    pr = c.post(f"/api/v1/customers/{hm_id}/devices", headers=M, json={"label": "HMIL gate HHT 1"}).json()
    raw = pr["pairing"].split("|", 1)[1]; pl_ = _json.loads(_b64.urlsafe_b64decode(raw + "=" * (-len(raw) % 4)))
    check(pr["pairing"].startswith("AISCUST1|") and pl_["c"] == "HMIL" and pl_["p"] == "CHN" and pl_["k"].startswith("aisd_") and isinstance(pl_["lan"], list)
          and "<svg" in pr["qr_svg"], f"pairing QR: customer, plant, device key, AIS-network addresses {pl_['lan']}, internet address '{pl_['wan']}'")
    DK = {"X-API-Key": pl_["k"], "X-Route": "WAN", "X-App-Version": "2.2.0"}
    r = c.post("/api/v1/users", headers=A, json={"user_id": "hmil.staff", "full_name": "HMIL Stores", "plant_code": "CHN", "customer_code": "HMIL",
                                                 "roles": ["CUST_STAFF"], "password": "Cust@1234"})
    check(r.status_code == 200, f"Customer Staff user created for HMIL: {r.text[:80]}")
    r = c.post("/api/v1/users", headers=A, json={"user_id": "bad.staff", "full_name": "x", "plant_code": "CHN", "roles": ["CUST_STAFF"], "password": "Cust@1234"})
    check(r.status_code == 400, "Customer Staff without a customer refused")
    r = c.post("/api/v1/auth/login", json={"user_id": "hmil.staff", "password": "Cust@1234", "device_id": "HHT-X", "app_version": "9.9.9"})
    check(r.status_code == 403 and "customer HHT" in r.text, "customer staff cannot sign in on an AIS HHT")
    r = c.post("/api/ext/v1/return-slip", headers=DK, json={"pallets": ["CHN-P002"], "vehicle_no": "TN09ZZ1"})
    check(r.status_code == 401 and "sign-in" in r.text, "customer HHT slip needs a staff sign-in")
    r = c.post("/api/ext/v1/staff-login", headers=DK, json={"user_id": "hmil.staff", "password": "wrong"})
    check(r.status_code == 401, "wrong staff password refused")
    st = c.post("/api/ext/v1/staff-login", headers=DK, json={"user_id": "HMIL.Staff", "password": "Cust@1234"}).json()
    check(st.get("token") and set(st["perms"]) == {"CUST_RECEIVE", "CUST_RETURN"}, f"staff signs in on the customer HHT: {st.get('perms')}")
    DK |= {"X-Staff-Token": st["token"]}
    me_ = c.get("/api/ext/v1/me", headers=DK).json()
    check(me_["customer"] == "HMIL" and me_["device"] == "HMIL gate HHT 1", "device key works on the customer API (/me)")
    v = c.post("/api/ext/v1/validate", headers=DK, json={"tags": held[:2] + ["NOPE-1"]}).json()
    check(v["ok_count"] == min(2, len(held)) and not v["results"][-1]["ok"], f"pallet check at the customer: {v['ok_count']} ok, unknown refused")
    sl = c.post("/api/ext/v1/return-slip", headers=DK, json={"pallets": held[:2], "vehicle_no": "tn 09 zz 1111", "customer_challan_no": "HM/DC/778", "client_ref": "dev-uuid-1"}).json()
    sl2 = c.post("/api/ext/v1/return-slip", headers=DK, json={"pallets": held[:2], "vehicle_no": "tn 09 zz 1111", "customer_challan_no": "HM/DC/778", "client_ref": "dev-uuid-1"}).json()
    with SessionLocal() as db:
        rs = db.get(models.ReturnSlip, sl["slip_no"]); dv = db.query(models.CustomerDevice).filter_by(customer_id=hm_id).first()
        src, by, route = rs.source, rs.created_by, dv.last_route
    check(sl["qr_payload"].startswith("AIS1|RS|") and sl2.get("duplicate") and src == "CUST_HHT" and by == "hmil.staff" and route == "WAN",
          f"Option A slip from the customer HHT: signed QR, resend not doubled, source CUST_HHT, made by the signed-in staff, route recorded ({route})")
    devs = c.get(f"/api/v1/customers/{hm_id}/devices", headers=M).json()
    check(len(devs) == 1 and devs[0]["last_route"] == "WAN" and devs[0]["app_version"] == "2.2.0", "paired devices listed with last route and app version")
    c.delete(f"/api/v1/customers/{hm_id}/devices/{devs[0]['id']}", headers=M)
    r = c.get("/api/ext/v1/me", headers=DK)
    check(r.status_code == 401 and "not paired" in r.json()["detail"], "removed HHT is refused at once")
    print("23. Option B: photo of the customer's manual document -> pallets read and matched, photo kept, yard compares")
    with SessionLocal() as db:
        for i, (t, n) in enumerate([("F", 5001), ("F", 5002), ("NF", 5003), ("RW", 5004), ("F", 15005)]):
            db.add(models.Pallet(pallet_no=f"AIS-CHN-{t}-{n:05d}-009{i:04d}", home_plant="CHN", status="AT_CUSTOMER", customer_code="HMIL", load_state="EMPTY"))
        db.commit()
    text = ("HYUNDAI MOTOR INDIA - RETURNABLE DC No 7781  Date 11-10-2026\nVehicle TN09AB4321\n1  F-5001\n2  F.5OO2\n3  NF 5003\n"
            "4  RW-5O04\n5  F 15006\n6  X-99881\nTotal 6 pallets")
    m = c.post("/api/v1/slips/doc-match", headers=H, json={"customer_code": "HMIL", "text": text}).json()
    mt = {x["ref"]: x["how"] for x in m["matched"]}; pb = {x["ref"]: x["how"] for x in m["probable"]}
    check(set(mt) == {"F-05001", "F-05002", "NF-05003", "RW-05004"} and pb.get("F-15005") == "one digit differs" and "X-99881" in " ".join(m["unread"]),
          f"document read: 4 matched (O->0 fixed), F 15006 -> probable F-15005, X-99881 unread; DC no / date ignored: {mt} {pb} {m['unread']}")
    rd = c.post("/api/v1/slips/doc-read", headers=H, data={"customer_code": "HMIL", "device_text": text}, files={"file": ("p.jpg", b"\xff\xd8x", "image/jpeg")}).json()
    check(rd["stated_total"] == 6 and len(rd["matched"]) == 4 and rd["ai"] is False, f"photo page read on the server: stated total {rd['stated_total']}, AI off -> HHT text used")
    import app.docread as _dr
    _old = _dr.ai_text; _dr.ai_text = lambda img, mt="image/jpeg": "F-5001\nRW 5004\nF-15005\nTOTAL=5"
    rd = c.post("/api/v1/slips/doc-read", headers=H, data={"customer_code": "HMIL", "device_text": "illegible scrawl"}, files={"file": ("p.jpg", b"\xff\xd8x", "image/jpeg")}).json()
    _dr.ai_text = _old
    check(rd["ai"] and rd["stated_total"] == 5 and {x["ref"] for x in rd["matched"]} == {"F-05001", "RW-05004", "F-15005"},
          f"AI reading (handwriting) merged and matched: {[x['ref'] for x in rd['matched']]}, total {rd['stated_total']}")
    docp = [x["pallet_no"] for x in m["matched"] + m["probable"]]
    r = c.post("/api/v1/slips", headers=H, json={"customer_code": "HMIL", "mode": "B2", "declared_qty": 6, "vehicle_no": "TN09AB4321",
                                                 "doc_pallets": docp, "doc_unread": m["unread"], "doc_text": text, "doc_read_qty": m["read"]}).json()
    bno = r["slip_no"]
    check(r["status"] == "PENDING_VERIFICATION" and r["declared_qty"] == 6 and r["doc_read_qty"] == 6, f"Option B slip keeps the count 6 and the document read ({r['warnings']})")
    up = c.post(f"/api/v1/slips/{bno}/doc-photo", headers=H, files={"file": ("doc.jpg", b"\xff\xd8\xff fake jpeg", "image/jpeg")}).json()
    ph = c.get(f"/api/v1/slips/{bno}/doc-photo/1", headers=P if False else H)
    check(up["photos"] == 1 and ph.status_code == 200 and ph.content.startswith(b"\xff\xd8"), "photo of the document kept on the server and viewable")
    for pn in ("AIS-CHN-F-05001-0090000", "AIS-CHN-NF-05003-0090002", "AIS-CHN-F-05002-0090001"):
        ev(c, H, "PALLET_SCAN_YARD", {"slip_no": bno, "scanned": pn})
    r2 = c.post("/api/v1/slips", headers=H, json={"customer_code": "HMIL", "mode": "B2", "declared_qty": 2, "vehicle_no": "TN09AB4322", "doc_text": "1 RW-5004  2 F 15005"}).json()
    with SessionLocal() as db:
        dp2 = db.get(models.ReturnSlip, r2["slip_no"]).doc_pallets
    check("AIS-CHN-RW-05004-0090003" in dp2 and "AIS-CHN-F-15005-0090004" in dp2, "slip made offline with only the photo text: matched on the server at sync")
    d = c.get(f"/api/v1/slips/{bno}", headers=H).json()["doc"]
    check(d["received"] == 3 and sorted(d["on_doc_not_received"]) == ["F-15005", "RW-05004"] and d["received_not_on_doc"] == [],
          f"yard receiving compared with the document: {d}")
    m = c.post("/api/v1/slips/doc-match", headers=H, json={"customer_code": "HMIL", "text": "Returnable DC 7790\nF - 2\nRW : 1 NOS\n"}).json()
    tc = {x["type"]: (x["doc_qty"], x["at_customer"]) for x in m["type_counts"]}
    check(m["doc_kind"] == "types" and tc.get("F", (0,))[0] == 2 and tc.get("RW", (0,))[0] == 1 and m["stated_total"] == 3 and not m["unread"],
          f"document with pallet types only (F - 2, RW : 1 NOS): read by type, total 3: {m['type_counts']}")
    m = c.post("/api/v1/slips/doc-match", headers=H, json={"customer_code": "HMIL", "text": "Empty pallets returned\nNo of pallets - 5\nDC 7791"}).json()
    check(m["doc_kind"] == "total" and m["stated_total"] == 5, f"document with the total only: 5 ({m['doc_kind']})")
    r3 = c.post("/api/v1/slips", headers=H, json={"customer_code": "HMIL", "mode": "B2", "declared_qty": 3, "vehicle_no": "TN09AB4323", "doc_text": "F - 2\nRW : 1 NOS"}).json()
    ev(c, H, "PALLET_SCAN_YARD", {"slip_no": r3["slip_no"], "scanned": "AIS-CHN-RW-05004-0090003"})
    d3 = c.get(f"/api/v1/slips/{r3['slip_no']}", headers=H).json()["doc"]
    bt = {x["type"]: (x["doc_qty"], x["received"]) for x in d3["by_type"]}
    check(d3["doc_total"] == 3 and bt.get("F") == (2, 0) and bt.get("RW") == (1, 1), f"yard compares by type: {d3['by_type']}")
    print("24. HHT setup QR; customer HHT without network: vehicle list, slip printed with the HHT's signed QR, opened at the gate, merged later")
    pg = c.get("/hht-setup")
    check(pg.status_code == 200 and "<svg" in pg.text and "Scan server QR" in pg.text, "server setup QR page (no login) for plant HHTs")
    with SessionLocal() as db:
        for i in range(3):
            db.add(models.Pallet(pallet_no=f"AIS-CHN-F-0600{i + 1}-009{1000 + i}", home_plant="CHN", status="AT_CUSTOMER", customer_code="HMIL", load_state="EMPTY"))
        db.commit()
    pr = c.post(f"/api/v1/customers/{hm_id}/devices", headers=M, json={"label": "HMIL gate HHT 2"}).json()
    raw = pr["pairing"].split("|", 1)[1]; pl2 = _json.loads(_b64.urlsafe_b64decode(raw + "=" * (-len(raw) % 4)))
    DK2 = {"X-API-Key": pl2["k"], "X-Route": "WAN"}
    DK2 |= {"X-Staff-Token": c.post("/api/ext/v1/staff-login", headers=DK2, json={"user_id": "hmil.staff", "password": "Cust@1234"}).json()["token"]}
    me2 = c.get("/api/ext/v1/me", headers=DK2).json()
    check(pl2.get("s") and pl2.get("i") and me2["sign_key"] == pl2["s"] and me2["device_id"] == pl2["i"] and "TN09AB4321" in me2["vehicles"],
          f"pairing QR and /me give the offline signing key and the vehicle list ({len(me2['vehicles'])} vehicles)")
    from app import services as _sv
    did, key = pl2["i"], pl2["s"]; ro = f"RO-CHN-{did}-00001"
    head = f"AIS1|RO|{ro}|CHN|HMIL|TN09AB7777|2|2610111200|{did}|F-06001,F-06002"
    qr = f"{head}|{_sv.offline_sig(key, head)}"
    bad = qr.replace("TN09AB7777", "TN09AB7778")
    g = c.post("/api/v1/slips/" + bad.replace("|", "%7C") + "/gate-in", headers=H)
    check(g.status_code == 400 and "check code" in g.text, "altered offline slip refused at the gate")
    g = c.post("/api/v1/slips/" + qr.replace("|", "%7C") + "/gate-in", headers=H, json={"vehicle_no": "TN09AB7777"})
    check(g.status_code == 200 and g.json()["slip_no"] == ro and g.json()["status"] == "IN_GATE", f"slip printed without network opened at the gate from its QR: {g.text[:120]}")
    with SessionLocal() as db:
        st = db.get(models.Pallet, "AIS-CHN-F-06001-0091000").status; nl = db.query(models.ReturnSlipLine).filter_by(slip_no=ro).count()
    check(st == "IN_RETURN" and nl == 2, f"pallets from the QR's type + serial list put IN_RETURN ({st}, {nl} lines)")
    r = c.post("/api/ext/v1/return-slip", headers=DK2, json={"slip_no": ro, "pallets": ["AIS-CHN-F-06001-0091000", "AIS-CHN-F-06002-0091001", "AIS-CHN-F-06003-0091002"],
                                                                 "vehicle_no": "TN09AB7777", "customer_challan_no": "HM/DC/901", "client_ref": "dev2-uuid-1"}).json()
    with SessionLocal() as db:
        nl = db.query(models.ReturnSlipLine).filter_by(slip_no=ro).count(); dc = db.get(models.ReturnSlip, ro).customer_challan_no
    check(r.get("duplicate") and nl == 3 and dc == "HM/DC/901" and r["slip_no"] == ro, "HHT online again: same slip number, details merged (3rd pallet, DC)")
    r = c.post("/api/ext/v1/return-slip", headers=DK2, json={"slip_no": "RO-CHN-999-00001", "pallets": ["AIS-CHN-F-06003-0091002"], "vehicle_no": "TN09AB7777", "customer_challan_no": "X1"})
    check(r.status_code == 400, "offline number of another HHT refused")
    r = c.post("/api/ext/v1/return-slip", headers=DK2, json={"slip_no": f"RO-CHN-{did}-00002", "pallets": ["AIS-CHN-F-06003-0091002"], "vehicle_no": "TN09AB7777"})
    r = r.json()
    check(r["slip_no"] == f"RO-CHN-{did}-00002" and r["declared_qty"] == 1, "offline slip without customer DC (optional) keeps its auto number")
    pu = c.get("/api/v1/sync/pull", headers=H).json()
    check("TN09AB7777" in pu.get("vehicles", []), f"plant HHT gets the vehicle list at sync ({len(pu.get('vehicles', []))})")
    print("25. Dispatch Bench trip and its GCS joined into one record")
    from app import bench as _bench, gcs as _gcs
    fake2 = {"/api/dispatches/recent?limit=300": {"dispatches": [
                {"id": 601, "trip_id": 91, "plant": "HVF1", "vehicle": "TN11Z1", "dispatch_status": "Completed", "sent_at": datetime.now().isoformat(), "pdi_by": "Ravi"},
                {"id": 602, "trip_id": 92, "plant": "HVF1", "vehicle": "TN11Z2", "dispatch_status": "Completed", "sent_at": datetime.now().isoformat()}]},
             "/api/picks?dispatch_id=601": {"picks": [{"picked_lpn": "BL61", "item_code": "WS-9", "qty": 20, "sub_inv": "HMIL-FG"}, {"picked_lpn": "BL62", "item_code": "WS-9", "qty": 20, "sub_inv": "HMIL-FG"}]},
             "/api/picks?dispatch_id=602": {"picks": [{"picked_lpn": "BL71", "item_code": "DR-3", "qty": 30, "sub_inv": "HMIL-FG"}]}}
    with SessionLocal() as db:
        _bench.pull(db, fetch=lambda path: fake2[path]); db.commit()
    rows = lambda cards: [{"lpn_no": f"F.6001.{x}", "invoice_no": "INV-9", "part_no": "WS-9", "part_qty": 20} for x in cards]
    with SessionLocal() as db:
        _gcs.apply_gcs(db, "CHN", {"gcs_no": "GCS/9001", "vehicle_no": "TN11Z1", "customer_code": "HMIL"}, "t", rows(["BL61", "BL62"])); db.commit()
        g = db.get(models.PickList, "GCS-CHN-GCS9001"); t = db.get(models.PickList, "DB-CHN-601"); l = db.get(models.Lpn, "BL61")
        st = (g.status, g.linked_picklist, t.status, t.linked_picklist, l.reserved_for, g.pdi_sign)
    check(st == ("OPEN", "DB-CHN-601", "LINKED", "GCS-CHN-GCS9001", "GCS-CHN-GCS9001", "Ravi"),
          f"trip not loaded yet: GCS sheet is the loading record, trip closed as LINKED, its part cards moved, PDI sign kept: {st}")
    e = ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": "DB-CHN-602", "scanned": "BL71", "pallet": "CHN-P005"})
    c.post("/api/v1/picklists/DB-CHN-602/confirm", headers=H)
    with SessionLocal() as db:
        _gcs.apply_gcs(db, "CHN", {"gcs_no": "GCS/9002", "vehicle_no": "TN11Z2", "customer_code": "HMIL"}, "t",
                       [{"lpn_no": "F.6002.BL71", "invoice_no": "INV-10", "part_no": "DR-3", "part_qty": 30}]); db.commit()
        g = db.get(models.PickList, "GCS-CHN-GCS9002"); t = db.get(models.PickList, "DB-CHN-602")
        st = (g.status, g.linked_picklist, t.gcs_no, t.invoice_no, db.query(models.PickListItem).filter_by(picklist_no="DB-CHN-602").count())
    check(st == ("LINKED", "DB-CHN-602", "GCS/9002", "INV-10", 1), f"trip already loaded: it takes the GCS no + invoice lines, GCS sheet LINKED (not loaded twice): {st}")
    o = c.post("/api/v1/picklists/gcs-open", headers=H, json={"gcs_no": "GCS/9002"}).json()
    check(o["picklist_no"] == "DB-CHN-602", "opening GCS 9002 on the HHT goes to the joined trip")
    print("26. Customer staff receive an AIS delivery by scanning each pallet while unloading")
    inc = c.get("/api/ext/v1/incoming", headers=DK2).json()["deliveries"]
    d0 = next((d for d in inc if d["expected"]), None)
    check(d0 is not None, f"incoming deliveries for HMIL listed with their pallets ({len(inc)})")
    pal0 = d0["pallets"][0]["pallet_no"]
    r = c.post("/api/ext/v1/receive", headers=DK2, json={"picklist_no": d0["picklist_no"], "scanned": [pal0, "AIS-CHN-F-06003-0091002"]}).json()
    check(r["received"] == 1 and r["results"][0]["on_load"] and not r["results"][1]["on_load"] and r["extra"],
          f"pallet on the delivery received; pallet not on it recorded as extra: {r['received']}/{r['expected']} extra {r['extra']}")
    r2 = c.post("/api/ext/v1/receive", headers=DK2, json={"picklist_no": d0["picklist_no"], "scanned": [pal0]}).json()
    check(r2["received"] == 1, "scanning the same pallet again is not counted twice")
    r3 = c.post("/api/ext/v1/receive", headers={k: v for k, v in DK2.items() if k != "X-Staff-Token"}, json={"picklist_no": d0["picklist_no"], "scanned": [pal0]})
    check(r3.status_code == 401, "receiving needs the staff sign-in")
    print("27. One server, many plants: plant chosen on the HHT, Bench trips go to their own plant")
    r = c.post("/api/v1/users", headers=A, json={"user_id": "multi", "full_name": "Two plants", "plant_code": "CHN", "extra_plants": "PUN;XXX", "roles": ["FGWH"], "password": "Abc@1234"})
    lg = c.post("/api/v1/auth/login", json={"user_id": "multi", "password": "Abc@1234"}).json()
    check([x["code"] for x in lg["plants"]] == ["CHN", "PUN"], f"login lists the plants the user may choose: {lg['plants']}")
    MH = {"Authorization": "Bearer " + lg["token"], "X-Device": "HHT", "X-Device-Id": "HHT-M"}
    check(c.get("/api/v1/sync/pull", headers=MH | {"X-Plant": "PUN"}).json()["plant"]["code"] == "PUN", "HHT works for the chosen plant (PUN): sync serves PUN data")
    check(c.get("/api/v1/sync/pull", headers=MH).json()["plant"]["code"] == "CHN", "without a choice: own plant")
    check(c.get("/api/v1/sync/pull", headers=MH | {"X-Plant": "TLJ"}).status_code == 403, "plant not allowed for the user refused")
    with SessionLocal() as db:
        from app import bench as _b2
        for k, v in (("bench_plant", None), ("bench_plant_map", "3001=CHN")):
            row_ = db.get(models.Setting, k) or models.Setting(key=k); row_.value = v; db.add(row_)
        db.commit()
        fake3 = {"/api/dispatches/recent?limit=300": {"dispatches": [
                    {"id": 701, "trip_id": 1, "customer_code": "HMIL", "ais_plant": "PUN", "vehicle": "MH12A1", "dispatch_status": "Completed", "sent_at": datetime.now().isoformat()},
                    {"id": 702, "trip_id": 2, "customer_code": "HMIL", "org_code": "3001", "vehicle": "TN01B2", "dispatch_status": "Completed", "sent_at": datetime.now().isoformat()},
                    {"id": 703, "trip_id": 3, "customer_code": "HMIL", "vehicle": "TN01B3", "dispatch_status": "Completed", "sent_at": datetime.now().isoformat()}]},
                 "/api/picks?dispatch_id=701": {"picks": []}, "/api/picks?dispatch_id=702": {"picks": []}, "/api/picks?dispatch_id=703": {"picks": []}}
        res = _b2.pull(db, fetch=lambda path: fake3[path]); db.commit()
        ok3 = db.get(models.PickList, "DB-PUN-701") is not None and db.get(models.PickList, "DB-CHN-702") is not None and db.query(models.PickList).filter(models.PickList.ext_ref == "703").count() == 0
    check(ok3 and res["no_plant"] == 1 and res["plants"] == {"PUN": 1, "CHN": 1},
          f"plant-wise Bench: trip 701 -> PUN, 702 (org 3001 mapped) -> CHN, 703 without plant not taken: {res}")
print(f"\nALL {ok_n} CHECKS PASSED")
os.remove(DB) if os.path.exists(DB) and not os.getenv("KEEP") else print("db kept:", DB)
