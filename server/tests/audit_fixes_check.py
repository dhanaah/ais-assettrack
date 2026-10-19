"""Checks for the fixes from the 13-Oct-2026 full review (server 1.23.0). Run:
    PYTHONPATH=. PALLET_BENCH_SCHEDULER=0 PALLET_PICKLIST_SOURCE=APP python3 tests/audit_fixes_check.py
Developed by DT"""
import os, sys, uuid, json, base64, threading, time
from datetime import datetime, timezone

sys.path.insert(0, os.path.dirname(__file__))
from testdb import fresh_db_url, bench_trip_for  # noqa: E402
os.environ["PALLET_DB_URL"] = fresh_db_url("audit")
os.environ["PALLET_INTEGRATION_MODE"] = "STUB"
os.environ["PALLET_BENCH_SCHEDULER"] = "0"
os.environ.setdefault("PALLET_PICKLIST_SOURCE", "APP")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from fastapi.testclient import TestClient
from app.main import app
from app import models, services, config
from app.db import SessionLocal
from app.security import hash_pw

FAILS = []


def check(ok, what):
    print(("  ok - " if ok else "FAIL: ") + what)
    if not ok:
        FAILS.append(what)


def now():
    return datetime.now(timezone.utc).replace(tzinfo=None).isoformat()


def setup():
    with SessionLocal() as db:
        db.add_all([models.Plant(code="CHN", name="Chennai", gstin="33AAA"), models.Plant(code="PUN", name="Pune", gstin="27AAA")]); db.flush()
        db.add(models.Customer(code="HMIL", name="Hyundai", plant_code="CHN"))
        db.add(models.Customer(code="FORD", name="Ford", plant_code="CHN"))
        for i in range(1, 8):
            db.add(models.Pallet(pallet_no=f"CHN-P{i:03d}", home_plant="CHN", location_plant="CHN", status="AVAILABLE", zone="FGWH"))
        db.add(models.Pallet(pallet_no="PUN-P900", home_plant="PUN", location_plant="PUN", status="AVAILABLE", zone="YARD"))
        for i in range(4):
            db.add(models.Pallet(pallet_no=f"AIS-CHN-F-0600{i+1}-009{1000+i}", home_plant="CHN", status="AT_CUSTOMER", customer_code="HMIL", load_state="EMPTY"))
        for uid, plant, roles, cust in (("chn", "CHN", ["FGWH", "QA", "LOG", "SEC_OUT", "SEC_IN", "YARD", "PACK", "PADMIN"], None),
                                        ("master", None, ["MASTER"], None), ("hmil.staff", "CHN", ["CUST_STAFF"], "HMIL")):
            u = models.User(user_id=uid, full_name=uid, plant_code=plant, password_hash=hash_pw("x"), must_change_pw=False,
                            supervisor_allowed=True, supervisor_pin_hash=hash_pw("1234"), customer_code=cust)
            u.roles = [models.UserRole(role_code=r) for r in roles]; db.add(u); db.flush()
        db.commit()


def tok(c, u, hht=True):
    r = c.post("/api/v1/auth/login", json={"user_id": u, "password": "x"}); assert r.status_code == 200, r.text
    h = {"Authorization": f"Bearer {r.json()['token']}"}
    return h | ({"X-Device": "HHT", "X-Device-Id": "HHT-TEST"} if hht else {})


def ev(c, h, t, payload, ts=None):
    r = c.post("/api/v1/sync/push", headers=h, json={"events": [{"event_id": str(uuid.uuid4()), "device_id": "HHT-T", "event_type": t, "payload": payload, "local_ts": ts or now()}]})
    assert r.status_code == 200, r.text
    return r.json()["results"][0]


def pal(no):
    with SessionLocal() as db:
        p = db.get(models.Pallet, no); return (p.status, p.customer_code, p.picklist_no)


with TestClient(app, raise_server_exceptions=False) as c:
    setup()
    H, M = tok(c, "chn"), tok(c, "master", hht=False)

    print("1. yard scan of a pallet on an open pick list takes it off the list - never dispatched twice")
    n1 = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "PALLET_ONLY", "customer_code": "HMIL", "qty": 1}).json()["picklist_no"]
    n2 = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "PALLET_ONLY", "customer_code": "FORD", "qty": 1}).json()["picklist_no"]
    ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": n1, "scanned": "CHN-P001"})
    r = ev(c, H, "PALLET_SCAN_YARD", {"scanned": "CHN-P001"})
    check("taken off pick list" in r["result"], f"yard says it was taken off {n1}: {r['result'][:90]}")
    with SessionLocal() as db:
        check(db.query(models.PickListLine).filter_by(picklist_no=n1, pallet_no="CHN-P001").count() == 0, "line removed from the first list")
    ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": n2, "scanned": "CHN-P001"})
    for n in (n2,):
        c.post(f"/api/v1/picklists/{n}/confirm", headers=H); c.post(f"/api/v1/picklists/{n}/challan", headers=H)
        c.post(f"/api/v1/picklists/{n}/approve", headers=H, json={"vehicle_no": "TN01X2"})
    with SessionLocal() as db:
        k = db.get(models.PickList, n2); sc = k.challan_no
    r = c.post(f"/api/v1/picklists/{n2}/gate-out", headers=H, json={"scanned": sc})
    check(r.status_code == 200 and pal("CHN-P001")[:2] == ("AT_CUSTOMER", "FORD"), "dispatched once, to FORD")
    r = c.post(f"/api/v1/picklists/{n2}/gate-out", headers=H, json={"scanned": sc})
    check(r.status_code == 400, f"second gate-out of the same load refused ({r.status_code})")

    print("2. gate-out refuses a load whose pallet is no longer on it")
    n3 = c.post("/api/v1/picklists", headers=H, json={"dispatch_type": "PALLET_ONLY", "customer_code": "HMIL", "qty": 1}).json()["picklist_no"]
    ev(c, H, "PALLET_SCAN_DOCK", {"picklist_no": n3, "scanned": "CHN-P002"})
    c.post(f"/api/v1/picklists/{n3}/confirm", headers=H); c.post(f"/api/v1/picklists/{n3}/challan", headers=H)
    c.post(f"/api/v1/picklists/{n3}/approve", headers=H, json={"vehicle_no": "TN01X3"})
    with SessionLocal() as db:
        p2 = db.get(models.Pallet, "CHN-P002"); p2.status = "AVAILABLE"; p2.picklist_no = None; db.commit()
        k = db.get(models.PickList, n3); sc = k.challan_no
    r = c.post(f"/api/v1/picklists/{n3}/gate-out", headers=H, json={"scanned": sc})
    check(r.status_code == 400 and "HOLD" in r.text, f"HOLD VEHICLE when a pallet left the load ({r.status_code})")

    print("3. offline slip sent after the slip was closed: closed slip unchanged")
    hm = None
    with SessionLocal() as db:
        hm = db.query(models.Customer).filter_by(code="HMIL").first().id
    pr = c.post(f"/api/v1/customers/{hm}/devices", headers=M, json={"label": "gate HHT"}).json()
    raw = pr["pairing"].split("|", 1)[1]; pl2 = json.loads(base64.urlsafe_b64decode(raw + "=" * (-len(raw) % 4)))
    DK = {"X-API-Key": pl2["k"], "X-Route": "LAN"}
    st = c.post("/api/ext/v1/staff-login", headers=DK, json={"user_id": "hmil.staff", "password": "x"}).json().get("token", "")
    DK |= {"X-Staff-Token": st}
    did, key = pl2["i"], pl2["s"]
    ro = f"RO-CHN-{did}-00001"
    head = f"AIS1|RO|{ro}|CHN|HMIL|TN09AB7777|1|2610111200|{did}|F-06001"
    qr = f"{head}|{services.offline_sig(key, head)}"
    g = c.post("/api/v1/slips/" + qr.replace("|", "%7C") + "/gate-in", headers=H, json={"vehicle_no": "TN09AB7777"})
    check(g.status_code == 200, f"offline slip opened at the gate ({g.status_code})")
    ev(c, H, "PALLET_SCAN_YARD", {"slip_no": ro, "scanned": "AIS-CHN-F-06001-0091000"})
    check(c.post(f"/api/v1/slips/{ro}/close", headers=H, json={}).status_code == 200, "closed at the yard")
    r = c.post("/api/ext/v1/return-slip", headers=DK, json={"slip_no": ro, "pallets": ["AIS-CHN-F-06001-0091000", "AIS-CHN-F-06002-0091001"], "vehicle_no": "TN09AB7777"})
    j = r.json()
    check(r.status_code == 200 and j["status"] == "CLOSED" and j["declared_qty"] == 1, "closed slip keeps its pallets and count")
    check(pal("AIS-CHN-F-06002-0091001")[0] == "AT_CUSTOMER", "extra pallet NOT moved to in-return")
    check(any("already closed" in w for w in j.get("warnings", [])), "warning about the extra pallet")

    print("4. same pallet twice on an offline slip: no server error")
    ro2 = f"RO-CHN-{did}-00002"
    head = f"AIS1|RO|{ro2}|CHN|HMIL|TN09AB7777|1|2610111200|{did}|"
    qr = f"{head}|{services.offline_sig(key, head)}"
    c.post("/api/v1/slips/" + qr.replace("|", "%7C") + "/gate-in", headers=H, json={"vehicle_no": "TN09AB7777"})
    r = c.post("/api/ext/v1/return-slip", headers=DK, json={"slip_no": ro2, "pallets": ["AIS-CHN-F-06003-0091002", "AIS-CHN-F-06003-0091002"], "vehicle_no": "TN09AB7777"})
    check(r.status_code == 200, f"merge with a duplicate pallet accepted ({r.status_code})")

    print("5. customer slip cannot move another plant's pallet")
    r = c.post("/api/ext/v1/return-slip", headers=DK, json={"pallets": ["PUN-P900"], "vehicle_no": "TN09AB1212"}).json()
    check(pal("PUN-P900")[0] == "AVAILABLE", "Pune pallet stays AVAILABLE")

    print("6. print pages escape what customers type")
    r = c.post("/api/ext/v1/return-slip", headers=DK, json={"pallets": ["AIS-CHN-F-06004-0091003"], "vehicle_no": "TN09AB3434",
                                                         "driver_name": "<script>alert(1)</script>", "client_ref": str(uuid.uuid4())}).json()
    t = c.post("/api/v1/auth/ticket", headers=H).json()["ticket"]
    page = c.get(f"/print/slip/{r['slip_no']}?tok={t}").text
    check(c.get(f"/print/slip/{r['slip_no']}?tok=" + H["Authorization"].split()[1]).status_code == 401, "sign-in token in a link is refused (2-minute ticket only)")
    check(c.get("/api/v1/reports/export/pallets.xlsx?tok=" + t).status_code == 200, "ticket opens an Excel download")
    check("<script>alert(1)</script>" not in page and "&lt;script&gt;" in page, "driver name shown as text, not run")

    print("7. Option B slip sent twice by the gate HHT: same slip back")
    body = {"customer_code": "HMIL", "mode": "B2", "declared_qty": 2, "doc_photo_count": 1, "vehicle_no": "TN09AB5555", "slip_no": "CHN-B-DEV-1"}
    r1 = c.post("/api/v1/slips", headers=H, json=body); r2 = c.post("/api/v1/slips", headers=H, json=body)
    check(r1.status_code == 200 and r2.status_code == 200 and r2.json().get("duplicate"), "repeat accepted as the same slip")

    print("8. Option B: pallet already in the yard does not count")
    ev(c, H, "PALLET_SCAN_YARD", {"slip_no": "CHN-B-DEV-1", "scanned": "CHN-P003"})
    with SessionLocal() as db:
        l = db.query(models.ReturnSlipLine).filter_by(slip_no="CHN-B-DEV-1", pallet_no="CHN-P003").first()
        check(l is not None and l.exception == "ALREADY_IN", "marked ALREADY_IN")

    print("9. document numbers never repeat when two requests run together")
    got = []
    def take():
        with SessionLocal() as db:
            got.append(services.doc_number(db, "CHN", "GE")); time.sleep(0.3); db.commit()
    th = [threading.Thread(target=take) for _ in range(4)]; [x.start() for x in th]; [x.join() for x in th]
    check(len(set(got)) == 4, f"4 different numbers {sorted(got)}")

    print("10. customer staff cannot sign in to the AIS system")
    r = c.post("/api/v1/auth/login", json={"user_id": "hmil.staff", "password": "x"})
    check(r.status_code == 403, f"refused ({r.status_code})")

    print("11. old public secret: new slips' label links / QRs made with it are refused")
    with SessionLocal() as db:
        sl = db.query(models.ReturnSlip).filter(models.ReturnSlip.mode == "A").order_by(models.ReturnSlip.created_at.desc()).first()
        no = sl.slip_no
    bad = services.label_token(no, config.LEGACY_SECRET)
    check(c.get(f"/api/ext/v1/return-slip/{no}/label?t={bad}").status_code == 403, "label link made with the old secret refused")
    check(c.get(f"/api/ext/v1/return-slip/{no}/label?t={services.label_token(no)}").status_code == 200, "real label link works")

    print("12. supervisor PIN guessing is stopped")
    from app.security import verify_supervisor_pin, _PIN_FAILS
    _PIN_FAILS.clear()
    codes = []
    with SessionLocal() as db:
        for i in range(10):
            try:
                verify_supervisor_pin(db, "CHN", f"{9000 + i}"); codes.append("no")
            except Exception as e:
                codes.append(getattr(e, "status_code", 0))
        check(429 in codes, f"paused after repeated wrong PINs {codes}")
        try:
            verify_supervisor_pin(db, "CHN", "1234"); check(False, "right PIN also paused during the lock")
        except Exception as e:
            check(getattr(e, "status_code", 0) == 429, "right PIN also paused during the lock")
    _PIN_FAILS.clear()

    print("13. scan older than the kept history is never re-applied")
    keep = config.EVENT_RETENTION_DAYS; config.EVENT_RETENTION_DAYS = 30
    r = ev(c, H, "PALLET_SCAN_YARD", {"scanned": "CHN-P004"}, ts="2026-02-01T00:00:00")
    check(r["status"] == "REJECTED" and "too old" in r["result"], "scan older than the kept history refused")
    r = ev(c, H, "PALLET_SCAN_YARD", {"scanned": "CHN-P005"}, ts="1970-01-01T00:10:00")
    check(r["status"] != "REJECTED", f"scan from an HHT with a reset clock (1970) still applied: {r['status']}")
    config.EVENT_RETENTION_DAYS = keep

    print("14. damage mark only for this plant's pallets")
    r = ev(c, H, "DAMAGE_MARK", {"scanned": "PUN-P900"})
    check(r["status"] == "REJECTED", "Pune pallet cannot be marked damaged at Chennai")

    print("15. Option A slip with a torn QR: supervisor PIN lets it in, typed number alone does not")
    r = c.post("/api/ext/v1/return-slip", headers=DK, json={"pallets": ["CHN-P006"], "vehicle_no": "TN09AB9898", "client_ref": str(uuid.uuid4())}).json()
    sno = r["slip_no"]
    r1 = c.post(f"/api/v1/slips/{sno}/gate-in", headers=H, json={})
    check(r1.status_code == 400 and "supervisor PIN" in r1.text, f"slip number without QR held ({r1.status_code})")
    r2 = c.post(f"/api/v1/slips/{sno}/gate-in", headers=H, json={"supervisor_pin": "1234"})
    check(r2.status_code == 200, f"let in with supervisor PIN ({r2.status_code} {r2.text[:80]})")

    print("16. retry job: key from job.key works; a call passed on by IIS / proxy without it is refused")
    check(c.post("/api/v1/jobs/retry", headers={"X-Forwarded-For": "10.1.2.3"}).status_code == 403, "proxied call without key refused")
    check(c.post("/api/v1/jobs/retry", headers={"X-Forwarded-For": "10.1.2.3", "X-Job-Key": config.JOB_KEY}).status_code == 200, "call with the job key accepted")

    print("17. password change ends older sign-ins")
    W2 = tok(c, "chn", hht=False)
    r = c.post("/api/v1/auth/change-password", headers=W2, json={"old_password": "x", "new_password": "newpass1"})
    check(r.status_code == 200 and r.json().get("token"), "changed, new token given")
    check(c.get("/api/v1/auth/me", headers=H).status_code == 401, "old token refused")
    check(c.get("/api/v1/auth/me", headers={"Authorization": "Bearer " + r.json()["token"]}).status_code == 200, "new token works")

    # after check 17 the chn password is newpass1 - fresh tokens for the runtime checks
    r = c.post("/api/v1/auth/login", json={"user_id": "chn", "password": "newpass1"}); assert r.status_code == 200, r.text
    HY = {"Authorization": f"Bearer {r.json()['token']}", "X-Device": "HHT", "X-Device-Id": "HHT-T"}
    print("18. runtime review (16-Oct): one bad event never breaks the push / jams the HHT queue")
    with SessionLocal() as db:
        db.add(models.ReturnSlip(slip_no="RS-RT1", plant_code="CHN", customer_code="HMIL", mode="B2", declared_qty=5, status="IN_GATE")); db.commit()
    r1 = ev(c, HY, "PALLET_SCAN_YARD", {"slip_no": "RS-RT1", "scanned": "NO-SUCH-TAG-1"})
    r2 = ev(c, HY, "PALLET_SCAN_YARD", {"slip_no": "RS-RT1", "scanned": "NO-SUCH-TAG-1"})
    check(r1["status"] == "EXCEPTION" and r2["status"] in ("EXCEPTION", "REJECTED"), f"unknown tag scanned twice on the same slip: no server error ({r1['status']}/{r2['status']})")
    with SessionLocal() as db:
        nrows = db.query(models.ReturnSlipLine).filter_by(slip_no="RS-RT1", pallet_no="NO-SUCH-TAG-1").count()
    check(nrows == 1, f"one quarantined line, not two ({nrows})")
    r = c.post("/api/v1/sync/push", headers=HY, json={"events": [
        {"event_id": str(uuid.uuid4()), "device_id": "HHT-T", "event_type": "HEARTBEAT", "payload": {"pending": "abc"}, "local_ts": now()},
        {"event_id": str(uuid.uuid4()), "device_id": "HHT-T", "event_type": "PALLET_MOVE", "payload": {"scanned": "CHN-P001", "to_zone": "FGWH"}, "local_ts": now() + "+05:30"}]})
    check(r.status_code == 200, f"mixed naive/aware timestamps and a bad heartbeat in one push: still 200 ({r.status_code} {r.text[:80]})")

    print("19. runtime review: smudged / foreign QR at the gates gives a clean message, not a server error")
    r = c.post("/api/v1/picklists/gcs-open", headers=HY, json={"gcs_no": "AIS1|GCS|G1|CHN|HMIL|TN01|abc|ts|sig"})
    check(r.status_code == 400, f"GCS QR with a bad number: 400 ({r.status_code})")
    r = c.post("/api/v1/slips/" + "AIS1|RS|RS-RT1|CHN|HMIL|TN01|abc|ts|-".replace("|", "%7C") + "/gate-in", headers=HY, json={})
    check(r.status_code in (200, 400), f"slip QR with a bad number: no server error ({r.status_code})")
    r = c.post("/api/v1/slips/RTS%7C/gate-in", headers=HY, json={})
    check(r.status_code in (400, 404), f"empty legacy slip QR: clean refusal ({r.status_code})")

    print("20. runtime review: corrupt upload file is refused with a message, blank Active keeps the value")
    MA = tok(c, "master", hht=False)
    r = c.post("/api/v1/masters/plants/upload", headers=MA, files={"file": ("x.xlsx", b"this is not excel")})
    check(r.status_code == 400 and "Not a valid" in r.text, f"masters upload garbage: clean 400 ({r.status_code} {r.text[:60]})")
    r = c.post("/api/v1/masters/customers/upload", headers=MA, files={"file": ("c.csv", "Code,Name (as on GCS),Plant,Active (Y/N)\nHMIL,Hyundai,CHN,\n".encode())})
    j = r.json()
    check(r.status_code == 200 and j["updated"] == 1 and not j["errors"], f"blank Active cell updates the row without error ({r.text[:90]})")
    with SessionLocal() as db:
        cst = db.query(models.Customer).filter_by(code="HMIL", plant_code="CHN").first()
        check(cst.active is True, "blank Active kept the stored value")

    print("21. GCS file plant name (column AC) is matched through the Plant master")
    from app import gcs as gcssvc
    with SessionLocal() as db:
        db.get(models.Plant, "CHN").gcs_plant_name = "Chennai"; db.commit()
    with SessionLocal() as db:
        what, pk = gcssvc.apply_gcs(db, "PUN", {"gcs_no": "GCS-AC1", "plant": "Chennai", "vehicle_no": "TN01AA1111", "customer_code": "HMIL"})
        db.commit()
        check(what == "created" and pk.plant_code == "CHN" and pk.picklist_no.startswith("GCS-CHN-"), f"file saying 'Chennai' lands on plant CHN even from another folder ({pk.plant_code})")
        _w2, pk2 = gcssvc.apply_gcs(db, "PUN", {"gcs_no": "GCS-AC2", "plant": "Mystery City", "vehicle_no": "TN01AA1112", "customer_code": "HMIL"})
        db.commit()
        check(pk2.plant_code == "PUN", f"unknown plant name keeps the folder plant ({pk2.plant_code})")
        _w3, pk3 = gcssvc.apply_gcs(db, "CHN", {"gcs_no": "GCS-AC3", "plant": "chennai", "vehicle_no": "TN01AA1113", "customer_code": "HMIL"})
        db.commit()
        check(pk3.plant_code == "CHN", "case does not matter")

    print("22. new pallet label first scanned at Vehicle loading: registered from the QR and linked to its GCS lines")
    import json as _json
    with SessionLocal() as db:
        db.add(models.PickList(picklist_no="CHN-GCS-RW", plant_code="CHN", customer_code="HMIL", dispatch_type="CUSTOMER", qty=1, status="OPEN",
                               gcs_no="GCS-RW", source="GCS", vehicle_no="TN87F1098")); db.flush()
        db.add(models.GcsLpn(picklist_no="CHN-GCS-RW", lpn_no="RW31LPN", lpn_label="RW.3100.P555", pallet_ref="RW.3100", pallet_no=None, part_no="WS-1", qty=5))
        db.add(models.GcsLpn(picklist_no="CHN-GCS-RW", lpn_no="RW32LPN", lpn_label="RW.3101.P556", pallet_ref="RW.3101", pallet_no=None, part_no="WS-1", qty=5))
        db.add(models.Lpn(lpn_no="RW31LPN", plant_code="CHN", part_no="WS-1", qty=5, status="RESERVED", reserved_for="CHN-GCS-RW")); db.commit()
        bench_trip_for(db, "CHN-GCS-RW")
    qr = _json.dumps({"uniquePalletID": "AIS-CHN-RW-3100-10001", "ownerPlant": "CHN", "palletType": "RW", "typeSerial": "3100", "globalSerial": "10001"})
    r = c.post("/api/v1/picklists/CHN-GCS-RW/load", headers=HY, json={"scanned": qr})
    check(r.status_code == 200 and r.json()["loaded"] == 1, f"new label registered and loaded against GCS line RW.3100 ({r.status_code} {r.text[:90]})")
    with SessionLocal() as db:
        g = db.query(models.GcsLpn).filter_by(picklist_no="CHN-GCS-RW").first()
        pal = db.get(models.Pallet, "AIS-CHN-RW-3100-10001")
        check(pal is not None and g.pallet_no == "AIS-CHN-RW-3100-10001", "pallet created and GCS line linked by type + serial")
    qr2 = _json.dumps({"uniquePalletID": "AIS-XXX-RW-9-10002", "ownerPlant": "XXX", "palletType": "RW"})
    r = c.post("/api/v1/picklists/CHN-GCS-RW/load", headers=HY, json={"scanned": qr2})
    check(r.status_code == 400 and "not in the Plant master" in r.text, "unknown owner plant still refused")

    with SessionLocal() as db:
        lst = [x for x in c.get("/api/v1/loading", headers=HY).json() if x["picklist_no"] == "CHN-GCS-RW"]
        check(lst and lst[0]["expected"] == 2 and lst[0]["loaded"] == 1, f"GCS naming 2 pallets (one never scanned yet) expects 2, not 1 ({lst[0]['expected'] if lst else '-'})")
    print("23. a queued GCS can be removed from the server (cancel) and its pallets are released")
    r = c.post("/api/v1/picklists/CHN-GCS-RW/cancel", headers=HY, json={})
    check(r.status_code == 200 and r.json()["status"] == "CANCELLED", f"queued GCS removed ({r.status_code} {r.text[:70]})")
    with SessionLocal() as db:
        pal = db.get(models.Pallet, "AIS-CHN-RW-3100-10001")
        check(pal.status == "AVAILABLE" and pal.picklist_no is None, f"its loaded pallet released ({pal.status})")
    r = c.get("/api/v1/loading", headers=HY).json()
    check(not any(x["picklist_no"] == "CHN-GCS-RW" for x in r), "removed GCS gone from the HHT loading list")

    print("24a. challan PDF: preview without the QR before loading is finished (it was refused before)")
    with SessionLocal() as db:
        db.add(models.PickList(picklist_no="CHN-GCS-CP", plant_code="CHN", customer_code="HMIL", dispatch_type="CUSTOMER", qty=1, status="OPEN",
                               gcs_no="GCS-CP", source="GCS", vehicle_no="TN87F3001")); db.commit()
    r = c.get("/api/v1/picklists/CHN-GCS-CP/challan.pdf", headers=HY)
    check(r.status_code == 200 and r.content[:4] == b"%PDF" and b"PREVIEW" in (r.headers.get("content-disposition") or "").encode(), f"challan preview opens ({r.status_code})")
    print("24. GCS PDF before Finish loading is a preview WITHOUT the gate QR; the real pass needs APPROVED")
    with SessionLocal() as db:
        db.add(models.PickList(picklist_no="CHN-GCS-PV", plant_code="CHN", customer_code="HMIL", dispatch_type="CUSTOMER", qty=2, status="OPEN",
                               gcs_no="GCS-PV", source="GCS", vehicle_no="TN87F2001")); db.commit()
    r = c.get("/api/v1/picklists/CHN-GCS-PV/gcs.pdf", headers=HY)
    check(r.status_code == 200 and r.content[:4] == b"%PDF" and b"PREVIEW" in (r.headers.get("content-disposition") or "").encode(), f"open GCS gives a PREVIEW pdf ({r.status_code} {r.headers.get('content-disposition')})")
    with SessionLocal() as db:
        k = db.get(models.PickList, "CHN-GCS-PV"); k.status = "APPROVED"; k.challan_no = "CH-PV"; db.commit()
    r2 = c.get("/api/v1/picklists/CHN-GCS-PV/gcs.pdf", headers=HY)
    check(r2.status_code == 200 and r2.content[:4] == b"%PDF" and b"PREVIEW" not in (r2.headers.get("content-disposition") or "").encode(), "finished GCS gives the real QR gate pass")
    check(len(r2.content) != len(r.content), "the two PDFs differ (QR present only on the finished one)")

    print("25. pallet registered (at the Yard) AFTER its GCS file came is accepted at loading; last pallet finishes by itself")
    with SessionLocal() as db:
        db.add(models.PickList(picklist_no="CHN-GCS-LT", plant_code="CHN", customer_code="HMIL", dispatch_type="CUSTOMER", qty=1, status="OPEN",
                               gcs_no="GCS-LT", source="GCS", vehicle_no="TN87F4001")); db.flush()
        db.add(models.GcsLpn(picklist_no="CHN-GCS-LT", lpn_no="LT1", lpn_label="LT.700.P1", pallet_ref="LT.700", pallet_no=None, part_no="WS-1", qty=5))
        db.add(models.Pallet(pallet_no="AIS-CHN-LT-700-20001", home_plant="CHN", location_plant="CHN", status="AVAILABLE", zone="FGWH", load_state="LOADED"))
        db.commit()
        bench_trip_for(db, "CHN-GCS-LT")
    r = c.post("/api/v1/picklists/CHN-GCS-LT/load", headers=HY, json={"scanned": "AIS-CHN-LT-700-20001"})
    j = r.json() if r.status_code == 200 else {}
    check(r.status_code == 200 and j.get("finished") and j.get("challan_no"), f"accepted, and the only pallet finished the loading (challan {j.get('challan_no')}) ({r.status_code} {r.text[:80]})")

    print("26. customer staff: a sign-in AIS no longer accepts is reported so the HHT asks to sign in again")
    import jwt as _jwt
    from datetime import timedelta as _td
    r = c.post("/api/ext/v1/staff-login", headers=DK, json={"user_id": "hmil.staff", "password": "x"})
    stok = r.json().get("token", "")
    me = c.get("/api/ext/v1/me", headers=DK | {"X-Staff-Token": stok}).json()
    check(r.status_code == 200 and me.get("staff") == "hmil.staff" and me.get("staff_ok") is True, "fresh sign-in accepted (/me staff_ok)")
    bad = _jwt.encode({"sub": "hmil.staff", "typ": "cust", "cust": "CHN/HMIL", "exp": models.utcnow() + _td(days=5)}, "another-server-secret", algorithm="HS256")
    me = c.get("/api/ext/v1/me", headers=DK | {"X-Staff-Token": bad}).json()
    check(me.get("staff_ok") is False, "sign-in made by another server (other secret): /me says staff_ok false")
    r = c.get("/api/ext/v1/holding", headers=DK | {"X-Staff-Token": bad})
    r2 = c.post("/api/ext/v1/return-slip", headers=DK | {"X-Staff-Token": bad}, json={"pallets": ["CHN-P001"], "vehicle_no": "TN01AB1234", "client_ref": "z9", "customer_challan_no": "DC9"})
    check(r2.status_code == 401 and r2.json()["detail"].startswith("STAFF_SIGNIN:"), f"staff action with a stale sign-in: 401 STAFF_SIGNIN ({r2.status_code} {r2.text[:80]})")
    with SessionLocal() as db:
        u = db.query(models.User).filter_by(user_id="hmil.staff").first(); u.password_hash = hash_pw("newpw123"); db.commit()
    me = c.get("/api/ext/v1/me", headers=DK | {"X-Staff-Token": stok}).json()
    check(me.get("staff_ok") is False, "password reset by AIS: the old sign-in is no longer accepted")

    print("27. approval review (18-Oct-2026): security")
    with SessionLocal() as db:
        db.get(models.Plant, "CHN").gcs_password = "FtpS3cret!"
        ua = models.User(user_id="uadmin", full_name="ua", plant_code=None, password_hash=hash_pw("x"), must_change_pw=False, perms_override="USER_ADMIN")
        ug = models.User(user_id="guard", full_name="g", plant_code="CHN", password_hash=hash_pw("x"), must_change_pw=False, supervisor_allowed=True)
        ug.roles = [models.UserRole(role_code="SEC_IN")]
        ca = models.User(user_id="cadm", full_name="ca", plant_code=None, password_hash=hash_pw("x"), must_change_pw=False)
        ca.roles = [models.UserRole(role_code="CADMIN")]
        db.add_all([ua, ug, ca])
        db.add(models.Pallet(pallet_no="PUN-P950", home_plant="PUN", location_plant="PUN", status="AVAILABLE", zone="YARD"))
        db.commit()
    UA, GD, CA = tok(c, "uadmin", hht=False), tok(c, "guard"), tok(c, "cadm", hht=False)
    pl = c.get("/api/v1/plants", headers=GD).json()
    check(all("gcs_password" not in x for x in pl) and any(x.get("has_password") for x in pl), "FTP password never sent to the web / HHT (has_password only)")
    r = c.post("/api/v1/users", headers=UA, json={"user_id": "uadmin", "full_name": "ua", "roles": ["CADMIN"]})
    check(r.status_code == 403, f"USER_ADMIN cannot make itself CADMIN ({r.status_code})")
    r = c.post("/api/v1/users", headers=UA, json={"user_id": "admin", "full_name": "Central Admin", "roles": ["CADMIN"], "password": "Hack3dPass"})
    check(r.status_code == 403, f"USER_ADMIN cannot reset the admin password ({r.status_code})")
    r = c.post("/api/v1/users", headers=CA, json={"user_id": "weak1", "full_name": "w", "plant_code": "CHN", "roles": ["YARD"], "password": "abc"})
    check(r.status_code == 400 and "8 characters" in r.text, f"weak password refused ({r.status_code} {r.text[:60]})")
    tk = c.post("/api/v1/auth/ticket", headers=CA).json()["ticket"]
    check(c.post(f"/api/v1/auth/ticket?tok={tk}").status_code == 401, "a link ticket cannot make a new ticket")
    check(c.get(f"/api/v1/auth/me?tok={tk}").status_code == 401, "a link ticket cannot be used as a sign-in")
    from app import config as _cfg
    nopv = __import__("jwt").encode({"sub": "chn", "uid": 1, "exp": models.utcnow() + __import__("datetime").timedelta(hours=1)}, _cfg.JWT_SECRET, algorithm="HS256")
    check(c.get("/api/v1/auth/me", headers={"Authorization": "Bearer " + nopv}).status_code == 401, "token without password version refused")
    r1 = c.post("/api/v1/auth/set-pin", headers=GD, json={"pin": "1234"})
    check(r1.status_code == 400 and "6 digits" in r1.text, "PIN needs 6+ digits")
    r1 = c.post("/api/v1/auth/set-pin", headers=GD, json={"pin": "111111"})
    check(r1.status_code == 400 and "simple" in r1.text, "too simple PIN refused")
    for _ in range(4):
        c.post("/api/v1/auth/set-pin", headers=GD, json={"pin": "123"})
    check(c.post("/api/v1/auth/set-pin", headers=GD, json={"pin": "583920"}).status_code == 429, "PIN changes are throttled (no guessing other supervisors' PINs)")
    r = c.get("/")
    check(r.headers.get("x-frame-options") and r.headers.get("x-content-type-options") == "nosniff" and "default-src 'self'" in (r.headers.get("content-security-policy") or ""), "security headers on every page")
    check(c.get("/docs").status_code == 404 and c.get("/openapi.json").status_code == 404, "Swagger / OpenAPI hidden unless PALLET_API_DOCS=1")
    from openpyxl import load_workbook
    import io as _io
    r = c.post("/api/v1/export/xlsx", headers=CA, json={"title": "t", "columns": ["a"], "rows": [["=HYPERLINK(\"http://x\",\"y\")"]]})
    ws = load_workbook(_io.BytesIO(r.content)).active
    vals = [cl for row in ws.iter_rows() for cl in row if isinstance(cl.value, str) and cl.value.startswith("=HYPER")]
    check(r.status_code == 200 and vals and all(cl.data_type == "s" for cl in vals), "Excel export keeps '=...' text as text (no formula)")
    o = c.get("/api/v1/pallets/PUN-P950", headers=GD).json()
    check(o.get("other_plant") and o["history"] == [], "another plant's pallet: no history shown")
    r = c.post("/api/v1/auth/login", json={"user_id": "admin", "password": config.BOOTSTRAP_ADMIN_PASS}, headers={"X-Forwarded-For": "1.2.3.4"})
    check(r.status_code in (200, 403), f"admin default password login handled ({r.status_code})")

    print("28. approval review (18-Oct-2026): logic")
    H = HY

    def _st(no):
        with SessionLocal() as db:
            return db.get(models.Pallet, no).status
    with SessionLocal() as db:
        db.add(models.Pallet(pallet_no="CHN-D001", home_plant="CHN", location_plant="CHN", status="DAMAGED", zone="YARD"))
        db.add(models.Pallet(pallet_no="CHN-X001", home_plant="CHN", status="AT_CUSTOMER", customer_code="HMIL"))
        db.add(models.ReturnSlip(slip_no="CHN-RTS-CL1", plant_code="CHN", customer_code="HMIL", mode="B2", status="CLOSED", declared_qty=1))
        db.commit()
    r = ev(c, H, "PALLET_SCAN_YARD", {"scanned": "CHN-D001"})
    check(_st("CHN-D001") == "DAMAGED" and "stays DAMAGED" in r["result"], f"damaged pallet stays DAMAGED at the Yard ({r['result'][:70]})")
    r = ev(c, H, "PALLET_SCAN_YARD", {"scanned": "CHN-X001", "slip_no": "CHN-RTS-CL1"})
    check(r["status"] == "REJECTED" and "CLOSED" in r["result"], "Yard scan on a CLOSED slip refused")
    r = c.post("/api/v1/pallets/CHN-X001/status", headers=H, json={"status": "AVAILABLE"})
    check(r.status_code == 400 and _st("CHN-X001") == "AT_CUSTOMER", "AT_CUSTOMER pallet cannot be set AVAILABLE by hand")
    r = c.post("/api/v1/slips/" + "PUN%2F00009" + "/gate-in", headers=H, json={"qr": "CHL|PUN/00009|X|1"})
    check(r.status_code != 404 or "Not Found" not in r.text, f"challan number with '/' reaches the gate-in route ({r.status_code} {r.text[:60]})")

    print("29. Dispatch Bench PDI check sheet <-> GCS: strict card match, reject + alternate, not invoiced, hard stops")
    from app import bench as _bn, gcs as _gc
    with SessionLocal() as db:
        for sn in (9101, 9102, 9103, 9104):
            db.add(models.Pallet(pallet_no=f"AIS-CHN-F-0{sn}-00{sn}0", home_plant="CHN", location_plant="CHN", status="AVAILABLE", zone="FGWH", load_state="LOADED", pallet_type="F"))
        db.commit()
    def bench_pull(trips):
        disp = {"dispatches": [{"id": i, "trip_id": i, "plant": "CHN", "vehicle": v, "dispatch_status": "Completed", "sent_at": now()} for i, v, _ in trips]}
        fk = {"/api/dispatches/recent?limit=300": disp} | {f"/api/picks?dispatch_id={i}": {"picks": [{"picked_lpn": cd, "item_code": "WS-1", "qty": 10} for cd in cards]} for i, _, cards in trips}
        with SessionLocal() as db:
            for k_, v_ in (("bench_plant", "CHN"),):
                row_ = db.get(models.Setting, k_) or models.Setting(key=k_); row_.value = v_; db.add(row_)
            db.commit()
            _bn.pull(db, fetch=lambda path: fk[path]); db.commit()
    def gcs_in(no, vehicle, cards):
        with SessionLocal() as db:
            _gc.apply_gcs(db, "CHN", {"gcs_no": no, "vehicle_no": vehicle, "customer_code": "HMIL"}, "t",
                          [{"lpn_no": f"F.{sn}.{cd}", "invoice_no": f"INV-{i % 2}", "part_no": "WS-1", "part_qty": 10} for i, (sn, cd) in enumerate(cards)]); db.commit()
        return "GCS-CHN-" + no.replace("/", "")
    # trip 901 on vehicle V1 carries A1 (pallet F.9101) and A2 (pallet F.9102)
    bench_pull([(901, "TN01V1", ["A1", "A2"])])
    g1 = gcs_in("G901", "TN01V1", [(9101, "A1"), (9102, "A2")])
    m = c.get(f"/api/v1/picklists/{g1}/pdi-match", headers=HY).json()
    check(m["trip_no"] == "DB-CHN-901" and m["ok"], f"every GCS card on the Bench trip: linked automatically, clean ({m['issues']})")
    # PDI reject of A2 in the Bench, alternate A3 on another pallet (F.9103) - Bench handles it, AssetTrack only sees the new list
    bench_pull([(901, "TN01V1", ["A1", "A3"])])
    m = c.get(f"/api/v1/picklists/{g1}/pdi-match", headers=HY).json()
    st = {x["lpn_no"]: x["status"] for x in m["cards"]}
    check(not m["ok"] and st.get("A2") == "REMOVED" and st.get("A3") == "NOT_INVOICED", f"after reject + alternate the old GCS no longer matches: {st}")
    r = c.post(f"/api/v1/picklists/{g1}/load", headers=HY, json={"scanned": "AIS-CHN-F-09102-0091020"})
    check(r.status_code == 400 and "removed from the trip by the Bench" in r.text, f"rejected card's pallet scanned at loading: STOP ({r.text[:90]})")
    with SessionLocal() as db:
        al = db.query(models.ScanMiss).filter_by(missed_point="PDI_CHECK", action="BLOCKED").count()
    check(al >= 1, "and an alert is recorded in the exception register")
    r = c.post(f"/api/v1/picklists/{g1}/load", headers=HY, json={"scanned": "AIS-CHN-F-09101-0091010"})
    check(r.status_code == 200 and not r.json().get("finished"), f"good card loads; finishing is held back ({r.json().get('finish_error', '')[:80]})")
    r = c.post(f"/api/v1/picklists/{g1}/confirm", headers=HY)
    check(r.status_code == 400 and "no challan" in r.text, f"finish / challan / QR refused while Bench and GCS differ: {r.text[:120]}")
    # GCS re-issued from the ERP with the alternate card
    gcs_in("G901", "TN01V1", [(9101, "A1"), (9103, "A3")])
    m = c.get(f"/api/v1/picklists/{g1}/pdi-match", headers=HY).json()
    check(m["ok"], f"re-issued GCS lists the alternate: clean again ({m['issues']})")
    r = c.post(f"/api/v1/picklists/{g1}/load", headers=HY, json={"scanned": "AIS-CHN-F-09103-0091030"}).json()
    check(r.get("finished") and r.get("challan_no"), f"alternate loaded: loading finished, challan + QR made ({r.get('challan_no')})")
    # not invoiced: Bench trip has a card the GCS does not list -> linked (GCS cards all on the trip) but blocked
    bench_pull([(902, "TN01V2", ["B1", "B2"])])
    g2 = gcs_in("G902", "TN01V2", [(9104, "B1")])
    m = c.get(f"/api/v1/picklists/{g2}/pdi-match", headers=HY).json()
    check(m["trip_no"] == "DB-CHN-902" and not m["ok"] and m["counts"]["not_invoiced"] == 1, f"PDI done but not invoiced: shown, blocks challan / QR ({m['issues']})")
    r = c.post(f"/api/v1/picklists/{g2}/load", headers=HY, json={"scanned": "AIS-CHN-F-09104-0091040"}).json()
    check(not r.get("finished") and r.get("pdi_ok") is False and "not invoiced" in (r.get("pdi_msg") or "") and not r.get("can_finish"),
          f"last pallet loaded but NOT finished, HHT shows why: {(r.get('pdi_msg') or '')[:90]}")
    # GCS card on no Bench trip -> no link, hard stop at loading
    g3 = gcs_in("G903", "TN01V3", [(9104, "Z9")])
    m = c.get(f"/api/v1/picklists/{g3}/pdi-match", headers=HY).json()
    check(m["trip_no"] is None and not m["ok"], "GCS card on no Bench trip: not linked")
    with SessionLocal() as db:
        db.get(models.Pallet, "AIS-CHN-F-09104-0091040").status = "AVAILABLE"; db.commit()
    r = c.post(f"/api/v1/picklists/{g3}/load", headers=HY, json={"scanned": "AIS-CHN-F-09104-0091040"})
    check(r.status_code == 400 and "HARD STOP" in r.text, f"loading against it: HARD STOP ({r.text[:90]})")
    # ambiguous: two open trips both carry every GCS card -> no automatic link
    bench_pull([(903, "TN01V4", ["D1"]), (904, "TN01V5", ["D1"])])
    g4 = gcs_in("G904", "TN01V4", [(9104, "D1")])
    with SessionLocal() as db:
        check(db.get(models.PickList, g4).linked_picklist is None, "two Bench trips carry the same card: not linked (needs a request)")
    # admin unlink of a wrong link
    r = c.post(f"/api/v1/picklists/{g2}/unlink", headers=CA, json={})
    check(r.status_code == 400 and "loaded" in r.text, "unlink refused once pallets are loaded")

    print("30. Dispatch Bench read from its SQL Server views (read only): OK / HOLD / REJECT, rework, API fallback")
    from sqlalchemy import create_engine, text as _t
    src = fresh_db_url("benchsrc"); se = create_engine(src)
    with se.begin() as x:
        x.execute(_t("CREATE TABLE src_trips (id INT, tno NVARCHAR(40), plant NVARCHAR(20), sec NVARCHAR(20), cust NVARCHAR(40), veh NVARCHAR(20), "
                     "st NVARCHAR(20), pdi NVARCHAR(80), planned DATETIME2, upd DATETIME2)"))
        x.execute(_t("CREATE TABLE src_cards (trip INT, lpn NVARCHAR(40), item NVARCHAR(40), qty INT, st NVARCHAR(10), why NVARCHAR(200), upd DATETIME2)"))
    with se.begin() as x:
        x.execute(_t("CREATE VIEW assettrack_trips AS SELECT CAST(id AS NVARCHAR(40)) AS trip_id, tno AS trip_no, plant AS plant_code, sec AS section, "
                     "cust AS customer_code, veh AS vehicle_no, st AS trip_status, pdi AS pdi_by, planned AS planned_at, upd AS updated_at FROM src_trips"))
    with se.begin() as x:
        x.execute(_t("CREATE VIEW assettrack_trip_cards AS SELECT CAST(trip AS NVARCHAR(40)) AS trip_id, lpn AS lpn_no, item AS part_no, qty, "
                     "st AS pdi_status, why AS pdi_reason, upd AS updated_at FROM src_cards"))

    def src_set(trips, cards):
        t0 = datetime.now()
        with se.begin() as x:
            x.execute(_t("DELETE FROM src_cards")); x.execute(_t("DELETE FROM src_trips"))
            for tid, veh in trips:
                x.execute(_t("INSERT INTO src_trips VALUES (:i, :n, 'CHN', 'MAIN', 'HMIL', :v, 'PDI_DONE', 'QA Ravi', :t, :t)"), {"i": tid, "n": f"T-{tid}", "v": veh, "t": t0})
            for tid, lpn, st_, why in cards:
                x.execute(_t("INSERT INTO src_cards VALUES (:i, :l, 'WS-1', 10, :s, :w, :t)"), {"i": tid, "l": lpn, "s": st_, "w": why, "t": t0})
    with SessionLocal() as db:
        for sn in (9201, 9202, 9203, 9204, 9205):
            db.add(models.Pallet(pallet_no=f"AIS-CHN-F-0{sn}-00{sn}0", home_plant="CHN", location_plant="CHN", status="AVAILABLE", zone="FGWH", load_state="LOADED", pallet_type="F"))
        for k_, v_ in (("bench_sql_url", src), ("bench_url", None), ("bench_user", None)):
            row_ = db.get(models.Setting, k_) or models.Setting(key=k_); row_.value = v_; db.add(row_)
        db.commit()
    r = c.post("/api/v1/integrations/dispatch-bench", headers=CA, json={"bench_enabled": False, "bench_source": "SQL", "bench_sql_database": "BenchDB",
                                                                       "bench_sql_trips_view": "assettrack_trips", "bench_sql_cards_view": "assettrack_trip_cards",
                                                                       "bench_plant": "CHN", "bench_days": 3})
    check(r.status_code == 200, f"Bench settings: read from SQL ({r.text[:80]})")
    r = c.post("/api/v1/integrations/dispatch-bench", headers=CA, json={"bench_source": "SQL", "bench_sql_database": "BenchDB", "bench_sql_trips_view": "x; DROP TABLE y"})
    check(r.status_code == 400, "a view name that is not a plain name is refused")
    src_set([(950, "TN01SQL")], [(950, "S1", "OK", None), (950, "S2", "HOLD", "edge chip - rework"), (950, "S3", "REJECT", "scratch"), (950, "S4", "OK", None)])
    r = c.post("/api/v1/integrations/dispatch-bench/test-sql", headers=CA).json()
    check(r.get("trips") == 1 and r.get("cards") == 4, f"Test SQL read: both views readable ({r})")
    r = c.post("/api/v1/integrations/dispatch-bench/pull", headers=CA).json()
    check(r.get("source") == "SQL" and r.get("created") == 1, f"pull from the SQL views creates the trip ({r})")
    with SessionLocal() as db:
        t = db.get(models.PickList, "DB-CHN-950")
        cs = {x.lpn_no: x for x in db.query(models.BenchTripCard).filter_by(trip_no="DB-CHN-950").all()}
        check(t is not None and t.vehicle_no == "TN01SQL" and "MAIN" in (t.remarks or "") and (t.pdi_sign or "").startswith("QA Ravi"), "trip: vehicle, section and PDI sign from the view")
        check(cs["S3"].removed_at is not None and cs["S3"].pdi_status == "REJECT" and cs["S3"].pdi_reason == "scratch", "REJECT: off the trip, kept with its reason")
        check(cs["S2"].removed_at is None and cs["S2"].pdi_status == "HOLD", "HOLD: stays on the trip")
        check(getattr(db.get(models.Lpn, "S3"), "reserved_for", None) != "DB-CHN-950" and db.get(models.Lpn, "S2").reserved_for == "DB-CHN-950", "rejected card not reserved; held card stays reserved")
        check("rejected at PDI - scratch" in (_gc.card_block(db, t, ["S3"]) or ""), "rejected card at loading: STOP with the reject reason")
    g5 = gcs_in("G950", "TN01SQL", [(9201, "S1"), (9202, "S2"), (9204, "S4")])
    m = c.get(f"/api/v1/picklists/{g5}/pdi-match", headers=HY).json()
    st = {x["lpn_no"]: x["status"] for x in m["cards"]}
    check(m["trip_no"] == "DB-CHN-950" and not m["ok"] and m["counts"]["hold"] == 1 and st.get("S2") == "HOLD" and st.get("S3") == "REMOVED",
          f"invoiced card on HOLD: linked but blocked ({m['issues']})")
    r = c.post(f"/api/v1/picklists/{g5}/load", headers=HY, json={"scanned": "AIS-CHN-F-09202-0092020"})
    check(r.status_code == 400 and "PDI HOLD" in r.text and "edge chip" in r.text, f"held card's pallet at loading: STOP ({r.text[:100]})")
    r = c.post(f"/api/v1/picklists/{g5}/load", headers=HY, json={"scanned": "AIS-CHN-F-09201-0092010"})
    check(r.status_code == 200, "PDI OK card loads")
    # rework done in the Bench: HOLD -> OK
    src_set([(950, "TN01SQL")], [(950, "S1", "OK", None), (950, "S2", "OK", None), (950, "S3", "REJECT", "scratch"), (950, "S4", "OK", None)])
    c.post("/api/v1/integrations/dispatch-bench/pull", headers=CA)
    m = c.get(f"/api/v1/picklists/{g5}/pdi-match", headers=HY).json()
    check(m["ok"], f"after rework (HOLD -> OK) Bench and GCS agree ({m['issues']})")
    c.post(f"/api/v1/picklists/{g5}/load", headers=HY, json={"scanned": "AIS-CHN-F-09202-0092020"})
    r = c.post(f"/api/v1/picklists/{g5}/load", headers=HY, json={"scanned": "AIS-CHN-F-09204-0092040"}).json()
    check(r.get("finished") and r.get("challan_no"), f"reworked card loaded: finished with challan ({r.get('challan_no')})")
    # card on HOLD and not invoiced: stays back, does not block the vehicle
    src_set([(951, "TN01SQL2")], [(951, "U1", "OK", None), (951, "U2", "HOLD", "rework")])
    c.post("/api/v1/integrations/dispatch-bench/pull", headers=CA)
    g6 = gcs_in("G951", "TN01SQL2", [(9205, "U1")])
    m = c.get(f"/api/v1/picklists/{g6}/pdi-match", headers=HY).json()
    st = {x["lpn_no"]: x["status"] for x in m["cards"]}
    check(m["trip_no"] == "DB-CHN-951" and m["ok"] and st.get("U2") == "HOLD_OUT", f"held card not invoiced: shown, not blocking ({st})")
    # SQL not readable + Bench API set up -> that pull uses the API
    fk = {"/api/dispatches/recent?limit=300": {"dispatches": [{"id": 952, "trip_id": 952, "plant": "CHN", "vehicle": "TN01API", "dispatch_status": "Completed", "sent_at": now()}]},
          "/api/picks?dispatch_id=952": {"picks": [{"picked_lpn": "W1", "item_code": "WS-1", "qty": 10}]}}
    real_get = _bn._get
    _bn._get = lambda cfg, path: fk[path]
    try:
        with SessionLocal() as db:
            for k_, v_ in (("bench_url", "http://bench:8765"), ("bench_user", "at"), ("bench_sql_trips_view", "no_such_view")):
                row_ = db.get(models.Setting, k_) or models.Setting(key=k_); row_.value = v_; db.add(row_)
            db.commit()
        r = c.post("/api/v1/integrations/dispatch-bench/pull", headers=CA).json()
        stt = c.get("/api/v1/integrations/dispatch-bench", headers=CA).json()
        check(r.get("source") == "API (SQL failed)" and "SQL read failed" in (stt["status"].get("last_error") or ""),
              f"SQL read fails: Bench API used, error shown ({r.get('source')})")
        check("bench_sql_url" not in stt["settings"], "the SQL connection string is never sent to the browser")
    finally:
        _bn._get = real_get
    se.dispose()

print()
print("ALL AUDIT CHECKS PASSED" if not FAILS else f"{len(FAILS)} FAILED")
sys.exit(1 if FAILS else 0)
