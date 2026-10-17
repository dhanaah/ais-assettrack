"""Master data in Excel: template with sample rows, current data download, upload (insert / update), delete.
Masters: plants, customers, transporters, pallet types, pallets. Sample rows have codes starting with SAMPLE and are
skipped on upload. Delete removes a record that was never used; a record already used in transactions is set
inactive (pallets: DISCONTINUED) so history stays correct. Developed by DT
"""
import csv, io
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from .. import models
from ..db import get_db
from ..security import need, current_user, Principal, audit

router = APIRouter(prefix="/api/v1/masters", tags=["masters-io"])

# column: (header, field, type, required, sample1, sample2)
SPECS = {
    "plants": {"model": models.Plant, "key": ["code"], "title": "Plants", "cols": [
        ("Code", "code", "str", True, "SAMPLE-A", "SAMPLE-B"), ("Name", "name", "str", True, "AIS Chennai", "AIS Pune (Varale)"),
        ("City", "city", "str", False, "Chennai", "Pune"), ("State", "state", "str", False, "Tamil Nadu", "Maharashtra"),
        ("GSTIN", "gstin", "str", False, "33AAACA0000A1Z1", "27AAACA0000A1Z1"),
        ("Address", "address", "str", False, "Plot F 76-81, Sipcot, Irungattukottai", "Gat 312, Varale, Talegaon"),
        ("Pallet Prefix", "pallet_prefix", "str", False, "CHN", "PUN"),
        ("Tag From", "tag_from", "int", False, 10001, 20001), ("Tag To", "tag_to", "int", False, 19999, 29999),
        ("Challan Source (APP/ORACLE)", "challan_source", "str", False, "APP", "ORACLE"),
        ("App Challan Series", "app_challan_series", "str", False, "CHN/", ""),
        ("Hold Days", "holding_days_default", "int", False, 30, 30),
        ("WMS Stale After (h)", "wms_stale_hours", "int", False, 4, 4), ("EBS Org Id", "ebs_org_id", "str", False, "", ""),
        ("Plant Name in GCS File", "gcs_plant_name", "str", False, "Chennai", "Pune"),
        ("GCS FTP Folder", "gcs_folder", "str", False, "/GCS/CHN", "/GCS/PUN"),
        ("GCS Printer", "gcs_printer", "str", False, "\\\\PRINTSRV\\FGWH-LASER", ""),
        ("GCS Print Mode (OFF/WINDOWS/RAW9100)", "gcs_print_mode", "str", False, "WINDOWS", "OFF"),
        ("GCS Copies", "gcs_copies", "int", False, 1, 1), ("Challan Copies (1-3)", "challan_copies", "int", False, 3, 3),
        ("Default Pallet HSN", "pallet_hsn", "str", False, "7326", "7326"),
        ("Active (Y/N)", "active", "bool", False, "Y", "Y")]},
    "customers": {"model": models.Customer, "key": ["code", "plant_code"], "title": "Customers", "cols": [
        ("Code", "code", "str", True, "SAMPLE-WEBASTO", "SAMPLE-NISSAN"),
        ("Name (as on GCS)", "name", "str", True, "WEBASTO ROOFSYSTEMS INDIA PRIVATE LIMITED", "RENAULT NISSAN AUTOMOTIVE INDIA PVT. LTD."),
        ("Plant", "plant_code", "str", True, "CHN", "CHN"), ("Type (OEM/OES/AFM)", "ctype", "str", False, "OEM", "OEM"),
        ("Ship To", "ship_to", "str", False, "Plot 12, SIPCOT", "Oragadam"), ("City", "city", "str", False, "Kancheepuram", "Oragadam"),
        ("State", "state", "str", False, "Tamil Nadu", "Tamil Nadu"), ("GSTIN", "gstin", "str", False, "33AAACW0000A1Z1", ""),
        ("Email To", "email_to", "str", False, "stores@customer.com", ""), ("Email CC", "email_cc", "str", False, "", ""),
        ("Contact Name", "contact_name", "str", False, "Mr. Kumar", ""), ("Contact Mobile", "contact_mobile", "str", False, "9840000000", ""),
        ("Hold Days", "holding_limit_days", "int", False, 30, 45), ("Remind Every (days)", "reminder_every_days", "int", False, 7, 7),
        ("Escalate After (days)", "escalation_days", "int", False, 45, 60), ("Chargeable (Y/N)", "chargeable", "bool", False, "Y", "N"),
        ("Return Mode (A/B)", "return_mode", "str", False, "B", "B"), ("Active (Y/N)", "active", "bool", False, "Y", "Y")]},
    "transporters": {"model": models.Transporter, "key": ["code"], "title": "Transporters", "cols": [
        ("Code", "code", "str", True, "SAMPLE-ARC", "SAMPLE-SRT"), ("Name", "name", "str", True, "ARC ENTERPRISES", "SRI RAGAVENDRA TRANSPORT"),
        ("GSTIN", "gstin", "str", False, "", ""), ("Contact Name", "contact_name", "str", False, "Mr. Raj", ""),
        ("Contact Mobile", "contact_mobile", "str", False, "9840000001", ""), ("Email", "email", "str", False, "", ""),
        ("Plants (; separated)", "plant_codes", "str", False, "CHN;PUN", "CHN"),
        ("Vehicles (; separated)", "vehicles", "str", False, "TN87D2115;TN31CE4390", "TN87B1498"), ("Active (Y/N)", "active", "bool", False, "Y", "Y")]},
    "pallet-types": {"model": models.PalletType, "key": ["code"], "title": "Pallet types", "cols": [
        ("Code", "code", "str", True, "SAMPLE-F", "SAMPLE-NF"), ("Description", "description", "str", False, "Steel returnable rack", "Steel rack (non-folding)"),
        ("HSN", "hsn", "str", False, "7326", "7326"), ("Value per Pallet (Rs)", "value", "float", False, 8500, 6200),
        ("Length (mm)", "length_mm", "int", False, 1600, 1500), ("Width (mm)", "width_mm", "int", False, 1150, 1100),
        ("Height (mm)", "height_mm", "int", False, 1300, 1250), ("Weight (kg)", "weight_kg", "float", False, 140, 110),
        ("Colour", "colour", "str", False, "Blue", "Grey"), ("Material", "material", "str", False, "MS steel", "MS steel"),
        ("Active (Y/N)", "active", "bool", False, "Y", "Y")]},
    "pallets": {"model": models.Pallet, "key": ["pallet_no"], "title": "Pallets", "cols": [
        ("Pallet No (QR id)", "pallet_no", "str", True, "SAMPLE-AIS-CHN-F-02948-0010001", "SAMPLE-AIS-CHN-NF-07944-0010003"),
        ("Home Plant", "home_plant", "str", True, "CHN", "CHN"), ("Pallet Type", "pallet_type", "str", False, "F", "NF"),
        ("Size", "size", "str", False, "1600x1150x1300", "1500x1100x1250"), ("Tare (kg)", "tare_kg", "float", False, 140, 110),
        ("Purchase Date (dd-mmm-yy)", "purchase_date", "date", False, "01-Apr-26", "15-May-26"), ("Notes", "notes", "str", False, "", "")]},
}
PERM = "GLOBAL_MASTERS"
# allowed values (upper case) per field - a wrong value is reported on its row, never saved
CHOICES = {"challan_source": ("APP", "ORACLE"), "gcs_print_mode": ("OFF", "WINDOWS", "RAW9100"), "challan_copies": (1, 2, 3)}


def _spec(kind):
    sp = SPECS.get(kind)
    if not sp:
        raise HTTPException(404, f"Unknown master '{kind}' - use one of {', '.join(SPECS)}")
    return sp


def _allow(p: Principal, kind: str):
    """Masters are central (GLOBAL_MASTERS); transporters may also be kept by plant logistics / plant admin
    (TRANSPORTER_MASTER) - for their own plant."""
    if p.has(PERM) or (kind == "transporters" and p.has("TRANSPORTER_MASTER")):
        return
    raise HTTPException(403, "Missing permission: " + ("TRANSPORTER_MASTER" if kind == "transporters" else PERM))


def _plant_user(p: Principal) -> str | None:
    return None if p.has(PERM) else p.plant


def _conv(v, typ):
    if v is None or (isinstance(v, str) and v.strip() in ("", "-")):
        return None
    if typ == "str":
        return (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)).strip()
    if typ == "int":
        return int(float(str(v).replace(",", "")))
    if typ == "float":
        return float(str(v).replace(",", ""))
    if typ == "bool":
        return str(v).strip().upper() in ("Y", "YES", "TRUE", "1", "ACTIVE")
    if typ == "date":
        if isinstance(v, datetime):
            return v
        for f in ("%d-%b-%y", "%d-%b-%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y"):
            try:
                return datetime.strptime(str(v).strip(), f)
            except ValueError:
                continue
        raise ValueError(f"date '{v}' - use dd-mmm-yy, e.g. 05-Oct-26")
    return v


def _xlsx(kind, rows, sample=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    sp = _spec(kind)
    wb = Workbook(); ws = wb.active; ws.title = sp["title"][:31]
    navy = PatternFill("solid", fgColor="1E3A8A"); req = PatternFill("solid", fgColor="C8102E")
    for c, col in enumerate(sp["cols"], 1):
        cell = ws.cell(row=1, column=c, value=col[0] + (" *" if col[3] else ""))
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = req if col[3] else navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = max(12, min(42, len(col[0]) + 4))
    ws.row_dimensions[1].height = 30; ws.freeze_panes = "A2"
    data = [[col[4] for col in sp["cols"]], [col[5] for col in sp["cols"]]] if sample else rows
    for r, vals in enumerate(data, 2):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, datetime):
                cell.number_format = "dd-mmm-yy"
            if sample:
                cell.font = Font(italic=True, color="64748B")
    if sample:
        hl = wb.create_sheet("How to fill")
        lines = [f"{sp['title']} - upload template (AIS AssetTrack)", "",
                 "1. Keep the first row (headings) exactly as it is. Red headings (*) are required.",
                 "2. The grey rows are SAMPLES (code starts with SAMPLE) - they are skipped on upload. Replace them with your data.",
                 "3. One row per record. A row whose code already exists UPDATES that record; a new code ADDS a record.",
                 "4. Dates: dd-mmm-yy (e.g. 05-Oct-26). Y/N columns: Y or N. Lists: separate with ;",
                 "5. Upload: web > Masters > " + sp["title"] + " > Upload. Errors are listed by row number; other rows are saved.", ""]
        lines += [f"{col[0]}{' (required)' if col[3] else ''}: example {col[4]!s}" for col in sp["cols"]]
        for i, t in enumerate(lines, 1):
            hl.cell(row=i, column=1, value=t).font = Font(bold=(i == 1))
        hl.column_dimensions["A"].width = 120
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def _send(buf, name, media="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"):
    return StreamingResponse(buf, media_type=media, headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/{kind}/template.xlsx")
def template(kind: str, p: Principal = Depends(current_user)):
    _allow(p, kind)
    return _send(_xlsx(kind, [], sample=True), f"AssetTrack_{_spec(kind)['title'].replace(' ', '_')}_Template.xlsx")


@router.get("/{kind}/data.xlsx")
def current(kind: str, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    _allow(p, kind)
    sp = _spec(kind)
    q = db.query(sp["model"])
    if kind == "transporters" and _plant_user(p):
        q = q.filter(sp["model"].plant_codes.like(f"%{p.plant}%"))
    if p.plant and hasattr(sp["model"], "plant_code"):
        q = q.filter(sp["model"].plant_code == p.plant)
    rows = []
    for o in q.all():
        r = []
        for col in sp["cols"]:
            v = getattr(o, col[1], None)
            r.append(("Y" if v else "N") if col[2] == "bool" else v)
        rows.append(r)
    return _send(_xlsx(kind, rows), f"AssetTrack_{sp['title'].replace(' ', '_')}_{datetime.now():%d-%b-%y}.xlsx")


def _read(name: str, data: bytes) -> list[dict]:
    if name.lower().endswith(".csv"):
        try:
            return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        except (UnicodeDecodeError, csv.Error):
            raise HTTPException(400, "Not a readable .csv file")
    from openpyxl import load_workbook
    try:
        ws = load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    except Exception:
        raise HTTPException(400, "Not a valid .xlsx file - download the template, fill it and upload that file")
    hdr = [str(c.value or "").replace("*", "").strip() for c in ws[1]]
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(x not in (None, "") for x in r)]


@router.post("/{kind}/upload")
async def upload(kind: str, request: Request, file: UploadFile = File(...), p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    _allow(p, kind)
    sp = _spec(kind); M = sp["model"]
    rows = _read(file.filename or "", await file.read())
    hmap = {col[0].lower(): col for col in sp["cols"]}
    out = {"inserted": 0, "updated": 0, "skipped": 0, "errors": []}
    for n, raw in enumerate(rows, 2):
        rec = {}
        sp_ = db.begin_nested()                     # each row on its own: a bad row never spoils the others
        try:
            for h, v in raw.items():
                col = hmap.get(str(h or "").replace("*", "").strip().lower())
                if col:
                    try:
                        rec[col[1]] = _conv(v, col[2])
                    except ValueError as ve:
                        raise ValueError(f"{col[0]}: '{v}' is not a {'number' if col[2] in ('int', 'float') else 'valid value'}"
                                         if col[2] != "date" else f"{col[0]}: {ve}")
            keyv = {k: rec.get(k) for k in sp["key"]}
            if any(str(v or "").upper().startswith("SAMPLE") for v in keyv.values()):
                sp_.rollback(); out["skipped"] += 1; continue
            miss = [col[0] for col in sp["cols"] if col[3] and rec.get(col[1]) in (None, "")]
            if miss:
                raise ValueError("missing " + ", ".join(miss))
            for k in ("code", "plant_code", "home_plant", "pallet_type", "challan_source", "gcs_print_mode"):
                if rec.get(k):
                    rec[k] = str(rec[k]).upper()
            for k, ok in CHOICES.items():
                if rec.get(k) not in (None, "") and rec[k] not in ok:
                    hdr = next(c[0] for c in sp["cols"] if c[1] == k) if any(c[1] == k for c in sp["cols"]) else k
                    raise ValueError(f"{hdr}: '{rec[k]}' - use {' / '.join(map(str, ok))}")
            keyv = {k: rec.get(k) for k in sp["key"]}
            if p.plant and "plant_code" in rec and rec["plant_code"] != p.plant:
                raise ValueError(f"plant {rec['plant_code']} - you can only load {p.plant}")
            for pk_field in ("plant_code", "home_plant"):
                if rec.get(pk_field) and not db.get(models.Plant, rec[pk_field]):
                    raise ValueError(f"plant {rec[pk_field]} not in the Plant master")
            obj = db.query(M).filter_by(**keyv).first()
            if kind == "transporters" and _plant_user(p):       # plant user: own plant always on it, other plants kept
                from .masters import plant_list_merge
                rec["plant_codes"] = plant_list_merge(obj.plant_codes if obj else None, rec.get("plant_codes"), p.plant)
            new = obj is None
            if new:
                obj = M(**keyv)
                if kind == "pallets":
                    obj.status, obj.location_plant = "AVAILABLE", rec.get("home_plant")
                db.add(obj)
            for k, v in rec.items():
                if k not in sp["key"] and v is not None:        # blank cell = keep the stored value (Active included)
                    setattr(obj, k, v)
            db.flush(); sp_.commit()
            out["inserted" if new else "updated"] += 1
        except Exception as e:
            sp_.rollback()
            out["errors"].append(f"row {n}: {str(e).splitlines()[0][:160]}")
    audit(db, p, "MASTER_UPLOAD", kind, file.filename, None, {k: v for k, v in out.items() if k != "errors"} | {"errors": len(out["errors"])}, request)
    db.commit()
    return out


def _used(db, kind, obj) -> str | None:
    """Why a record cannot be removed (it appears in transactions) - None when it is safe to delete."""
    if kind == "plants":
        if db.query(models.PickList).filter_by(plant_code=obj.code).first() or db.query(models.Pallet).filter_by(home_plant=obj.code).first():
            return "has pallets / pick lists"
    if kind == "customers":
        if db.query(models.PickList).filter_by(plant_code=obj.plant_code, customer_code=obj.code).first():
            return "has pick lists / gate passes"
    if kind == "transporters":
        if db.query(models.PickList).filter_by(transporter_code=obj.code).first():
            return "has pick lists / gate passes"
    if kind == "pallet-types":
        if db.query(models.Pallet).filter_by(pallet_type=obj.code).first():
            return "pallets of this type exist"
    if kind == "pallets":
        if db.query(models.PickListLine).filter_by(pallet_no=obj.pallet_no).first() or db.query(models.ReturnSlipLine).filter_by(pallet_no=obj.pallet_no).first():
            return "was dispatched / returned"
    return None


@router.delete("/{kind}/{key}")
def delete(kind: str, key: str, request: Request, plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    _allow(p, kind)
    sp = _spec(kind); M = sp["model"]
    if kind == "customers":
        obj = db.get(M, int(key)) if key.isdigit() else db.query(M).filter_by(code=key.upper(), plant_code=(plant or p.plant or "").upper()).first()
    else:
        obj = db.get(M, key if kind == "pallets" else key.upper())
    if not obj:
        raise HTTPException(404, "Not found")
    if getattr(obj, "plant_code", None) and p.plant and obj.plant_code != p.plant:
        raise HTTPException(403, "Other plant's record")
    if kind == "transporters" and _plant_user(p):
        pls = [x for x in (obj.plant_codes or "").split(";") if x]
        if p.plant not in pls:
            raise HTTPException(403, "Transporter not used by your plant")
        if len(pls) > 1:                                  # shared with other plants: only remove it from this plant
            obj.plant_codes = ";".join(x for x in pls if x != p.plant)
            audit(db, p, "MASTER_DELETE", kind, key, None, {"removed_plant": p.plant}, request); db.commit()
            return {"message": f"Transporter {obj.code} removed from {p.plant} (other plants still use it)"}
    why = _used(db, kind, obj)
    label = getattr(obj, "code", None) or getattr(obj, "pallet_no", key)
    if why:
        if kind == "pallets":
            if obj.status in ("ALLOCATED", "AT_CUSTOMER", "IN_RETURN"):
                raise HTTPException(400, f"Pallet {label} is {obj.status} - cannot be removed now")
            obj.status, obj.discontinued_at = "DISCONTINUED", datetime.utcnow()
            res = f"Pallet {label} {why} - marked DISCONTINUED (history kept)"
        else:
            obj.active = False
            res = f"{sp['title'][:-1]} {label} {why} - set INACTIVE (history kept)"
        action = "deactivated"
    else:
        if kind == "pallets":
            db.query(models.PalletHistory).filter_by(pallet_no=obj.pallet_no).delete()
        db.delete(obj); res = f"{sp['title'][:-1]} {label} deleted"; action = "deleted"
    audit(db, p, "MASTER_DELETE", kind, str(label), None, {"result": action}, request)
    db.commit()
    return {"result": action, "message": res}
