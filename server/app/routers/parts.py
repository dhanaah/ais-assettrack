"""FG parts: blanket orders, sub-inventory rules, WMS stock (Excel + EBS pull), LPNs & reservations,
PDI, EBS sub-inventory transactions, in-plant movement, other-plant receipt, missed scans, loaded/empty view.
Developed by DT"""
import csv, io
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session
from .. import models, services, lpn as lpnsvc
from ..db import get_db
from ..security import current_user, need, Principal, audit, verify_supervisor_pin
from ..models import utcnow

router = APIRouter(prefix="/api/v1", tags=["parts"])


def row(o):
    return {c.name: getattr(o, c.name) for c in o.__table__.columns}


def _plant(p: Principal, plant: str | None) -> str:
    pl = p.plant or plant
    if not pl:
        raise HTTPException(400, "plant required")
    p.require_plant(pl)
    return pl


def _rule(fn):
    try:
        return fn()
    except services.RuleError as ex:
        raise HTTPException(400, str(ex))


# ---------------------------------------------------------------- blanket orders
class BlanketIn(BaseModel):
    id: int | None = None
    plant_code: str
    customer_code: str
    part_no: str
    customer_part_no: str | None = None
    blanket_no: str
    po_number: str
    schedule_qty: int
    valid_from: datetime | None = None
    valid_to: datetime | None = None
    dispatch_type: str = "CUSTOMER"
    to_plant: str | None = None
    active: bool = True
    notes: str | None = None


@router.get("/blankets")
def list_blankets(plant: str | None = None, customer: str | None = None, part: str | None = None, active_only: bool = False,
                  p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(models.BlanketOrder)
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter_by(plant_code=pl)
    if customer:
        q = q.filter_by(customer_code=customer)
    if part:
        q = q.filter(models.BlanketOrder.part_no.like(f"%{part}%"))
    if active_only:
        q = q.filter_by(active=True)
    return [row(b) | {"open_qty": lpnsvc.blanket_open_qty(db, b),
                      "free_stock": int(db.query(func.coalesce(func.sum(models.Lpn.qty), 0)).filter(
                          models.Lpn.plant_code == b.plant_code, models.Lpn.part_no == b.part_no, models.Lpn.status == "AVAILABLE").scalar() or 0)}
            for b in q.order_by(models.BlanketOrder.customer_code, models.BlanketOrder.part_no).limit(2000).all()]


@router.post("/blankets")
def upsert_blanket(body: BlanketIn, request: Request, p: Principal = Depends(need("BLANKET_MASTER")), db: Session = Depends(get_db)):
    p.require_plant(body.plant_code)
    dt = body.dispatch_type.upper()
    if dt not in lpnsvc.PART_TYPES:
        raise HTTPException(400, "dispatch_type CUSTOMER or STOCK_TRANSFER")
    if dt == "STOCK_TRANSFER" and not body.to_plant:
        raise HTTPException(400, "to_plant required for stock transfer")
    b = db.get(models.BlanketOrder, body.id) if body.id else db.query(models.BlanketOrder).filter_by(
        plant_code=body.plant_code, customer_code=body.customer_code, part_no=body.part_no, po_number=body.po_number).first()
    old = row(b) if b else None
    if not b:
        b = models.BlanketOrder(); db.add(b)
    for k, v in body.model_dump(exclude={"id"}).items():
        setattr(b, k, v.upper() if k == "dispatch_type" else v)
    if b.schedule_qty < (b.released_qty or 0):
        raise HTTPException(400, f"Schedule qty below already released {b.released_qty}")
    audit(db, p, "BLANKET_SAVE", "blanket", f"{b.blanket_no}/{b.part_no}", old, body.model_dump(), request)
    db.commit()
    return row(b)


# ---------------------------------------------------------------- sub-inventory rules
class SubInvIn(BaseModel):
    id: int | None = None
    plant_code: str
    customer_code: str | None = None
    part_no: str | None = None
    ok_subinv: str
    ok_locator: str | None = None
    reject_subinv: str | None = None
    active: bool = True


@router.get("/subinv-rules")
def list_rules(plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(models.SubInvRule)
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter_by(plant_code=pl)
    return [row(r) for r in q.all()]


@router.post("/subinv-rules")
def upsert_rule(body: SubInvIn, request: Request, p: Principal = Depends(need("BLANKET_MASTER")), db: Session = Depends(get_db)):
    p.require_plant(body.plant_code)
    r = db.get(models.SubInvRule, body.id) if body.id else None
    if not r:
        r = models.SubInvRule(); db.add(r)
    for k, v in body.model_dump(exclude={"id"}).items():
        setattr(r, k, v or None if k in ("customer_code", "part_no", "ok_locator", "reject_subinv") else v)
    audit(db, p, "SUBINV_RULE_SAVE", "subinv_rule", str(body.id or "new"), None, body.model_dump(), request)
    db.commit()
    return row(r)


# ---------------------------------------------------------------- WMS stock
@router.get("/wms/status")
def wms_status(plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    if p.plant or plant:
        return lpnsvc.wms_status(db, _plant(p, plant))
    return [lpnsvc.wms_status(db, x.code) for x in db.query(models.Plant).filter_by(active=True).all()]


@router.get("/wms/uploads")
def wms_uploads(plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(models.WmsUpload).order_by(models.WmsUpload.ts.desc())
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter_by(plant_code=pl)
    return [row(x) for x in q.limit(100).all()]


@router.get("/wms/template")
def wms_template():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["LPN", "PART_NO", "PART_DESC", "QTY", "SUBINVENTORY", "LOCATOR", "PALLET_NO"])
    w.writerow(["LPN0001234", "WS-1234-FR", "Windshield front", 40, "FG-STORE", "A-01-01", "CHN-P00012"])
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": "attachment; filename=WMS_Stock_Upload_Template.csv"})


def _read_rows(name: str, data: bytes) -> list[dict]:
    if name.lower().endswith(".csv"):
        txt = data.decode("utf-8-sig", errors="ignore")
        return list(csv.DictReader(io.StringIO(txt)))
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = None
    out = []
    for r in it:
        if hdr is None:
            if r and any(str(c or "").strip().upper().replace(" ", "_") in ("LPN", "LPN_NO", "LICENSE_PLATE_NUMBER") for c in r):
                hdr = [str(c or "").strip() for c in r]
            continue
        if any(c not in (None, "") for c in r):
            out.append(dict(zip(hdr, r)))
    if hdr is None:
        raise HTTPException(400, "Header row with LPN column not found")
    return out


@router.post("/wms/upload")
async def wms_upload(request: Request, plant: str | None = Query(None), file: UploadFile = File(...),
                     p: Principal = Depends(need("WMS_UPLOAD")), db: Session = Depends(get_db)):
    pl = _plant(p, plant)
    rows = _read_rows(file.filename or "x.xlsx", await file.read())
    res = _rule(lambda: lpnsvc.apply_wms_snapshot(db, pl, rows, "EXCEL", p.user_id))
    audit(db, p, "WMS_UPLOAD", "wms", res["batch_id"], None, {k: v for k, v in res.items() if k != "warnings"}, request)
    db.commit()
    return res


@router.post("/wms/pull")
def wms_pull(request: Request, plant: str | None = None, p: Principal = Depends(need("WMS_UPLOAD")), db: Session = Depends(get_db)):
    from .. import integration
    pl = _plant(p, plant)
    ok, rows, msg = integration.pull_wms_stock(db, pl)
    if not ok:
        db.commit()
        raise HTTPException(502, f"EBS WMS pull failed: {msg}")
    if not rows:
        db.commit()
        raise HTTPException(400, "EBS returned no stock rows - nothing changed (upload Excel if needed)")
    res = _rule(lambda: lpnsvc.apply_wms_snapshot(db, pl, rows, "API", p.user_id))
    audit(db, p, "WMS_PULL", "wms", res["batch_id"], None, {k: v for k, v in res.items() if k != "warnings"}, request)
    db.commit()
    return res


# ---------------------------------------------------------------- LPNs & reservations
@router.get("/lpns")
def list_lpns(plant: str | None = None, status: str | None = None, part: str | None = None, pallet: str | None = None,
              picklist: str | None = None, flagged: bool = False, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(models.Lpn)
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter_by(plant_code=pl)
    if status:
        q = q.filter(models.Lpn.status.in_(status.split(",")))
    if part:
        q = q.filter(models.Lpn.part_no.like(f"%{part}%"))
    if pallet:
        q = q.filter_by(pallet_no=pallet)
    if picklist:
        q = q.filter((models.Lpn.picklist_no == picklist) | (models.Lpn.reserved_for == picklist))
    if flagged:
        q = q.filter(models.Lpn.wms_missing.is_(True))
    out = []
    now = utcnow()
    for l in q.order_by(models.Lpn.updated_at.desc()).limit(3000).all():
        d = row(l)
        d["reserved_hours"] = round((now - l.reserved_at).total_seconds() / 3600, 1) if l.reserved_at and l.status == "RESERVED" else None
        out.append(d)
    return out


@router.get("/lpns/{lpn_no}")
def get_lpn(lpn_no: str, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    l = db.get(models.Lpn, lpn_no)
    if not l:
        raise HTTPException(404, "LPN not found")
    p.require_plant(l.plant_code)
    return {"lpn": row(l), "pallet": row(db.get(models.Pallet, l.pallet_no)) if l.pallet_no and db.get(models.Pallet, l.pallet_no) else None,
            "txns": [row(t) for t in db.query(models.LpnTxn).filter_by(lpn_no=lpn_no).order_by(models.LpnTxn.created_at.desc()).all()],
            "picklists": [row(r) for r in db.query(models.PickListLpn).filter_by(lpn_no=lpn_no).all()]}


@router.get("/reservations")
def reservations(plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    """Reserved-LPN control: per pick list, reserved vs picked, age and WMS flags."""
    q = db.query(models.Lpn).filter(models.Lpn.reserved_for.isnot(None), models.Lpn.status.in_(["RESERVED", "PICKED", "PDI_OK"]))
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter_by(plant_code=pl)
    by = {}
    now = utcnow()
    for l in q.all():
        g = by.setdefault(l.reserved_for, {"picklist_no": l.reserved_for, "part_no": l.part_no, "reserved_lpn": 0, "reserved_qty": 0,
                                           "picked_lpn": 0, "picked_qty": 0, "oldest_hours": 0, "wms_missing": 0})
        if l.status == "RESERVED":
            g["reserved_lpn"] += 1; g["reserved_qty"] += l.qty
        else:
            g["picked_lpn"] += 1; g["picked_qty"] += l.qty
        if l.reserved_at:
            g["oldest_hours"] = max(g["oldest_hours"], round((now - l.reserved_at).total_seconds() / 3600, 1))
        g["wms_missing"] += 1 if l.wms_missing else 0
    for g in by.values():
        k = db.get(models.PickList, g["picklist_no"])
        g.update({"status": k.status if k else "?", "customer": k.customer_code if k else None, "part_qty": k.part_qty if k else None})
    return sorted(by.values(), key=lambda x: -x["oldest_hours"])


class ReleaseIn(BaseModel):
    supervisor_pin: str
    reason: str


@router.post("/lpns/{lpn_no}/release")
def release_lpn(lpn_no: str, body: ReleaseIn, request: Request, p: Principal = Depends(need("PICKLIST_CREATE")), db: Session = Depends(get_db)):
    l = db.get(models.Lpn, lpn_no)
    if not l:
        raise HTTPException(404, "LPN not found")
    p.require_plant(l.plant_code)
    if l.status != "RESERVED":
        raise HTTPException(400, f"LPN is {l.status} - only RESERVED can be released (un-pick at dock first)")
    if not verify_supervisor_pin(db, l.plant_code, body.supervisor_pin):
        raise HTTPException(403, "Supervisor PIN invalid")
    was = l.reserved_for
    l.status, l.reserved_for, l.reserved_at = ("MISSING" if l.wms_missing else "AVAILABLE"), None, None
    audit(db, p, "LPN_RESERVATION_RELEASE", "lpn", lpn_no, {"reserved_for": was}, {"reason": body.reason}, request)
    db.commit()
    return row(l)


# ---------------------------------------------------------------- PDI
class PdiIn(BaseModel):
    scanned: str
    result: str = "OK"
    remarks: str | None = None


@router.get("/picklists/{no}/lpns")
def picklist_lpns(no: str, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    k = db.get(models.PickList, no)
    if not k:
        raise HTTPException(404, "Not found")
    p.require_plant(k.plant_code)
    rows_ = db.query(models.PickListLpn).filter_by(picklist_no=no).all()
    txns = db.query(models.LpnTxn).filter_by(picklist_no=no).order_by(models.LpnTxn.created_at.desc()).all()
    reserved = db.query(models.Lpn).filter_by(reserved_for=no, status="RESERVED").all()
    return {"picklist": row(k), "lpns": [row(r) for r in rows_], "reserved": [row(x) for x in reserved], "txns": [row(t) for t in txns],
            "picked_qty": lpnsvc.picked_qty(db, no)}


@router.post("/picklists/{no}/pdi")
def pdi_mark(no: str, body: PdiIn, request: Request, p: Principal = Depends(need("PDI_CHECK")), db: Session = Depends(get_db)):
    k = db.get(models.PickList, no)
    if not k:
        raise HTTPException(404, "Not found")
    p.require_plant(k.plant_code)
    try:
        msg, alerts = lpnsvc.pdi_mark(db, k, body.scanned, body.result, body.remarks, p.user_id)
    except services.RuleError as ex:
        db.commit()       # keep the missed-scan record
        raise HTTPException(400, str(ex))
    audit(db, p, "PDI_MARK", "picklist", no, None, body.model_dump(), request)
    db.commit()
    return {"message": msg, "alerts": alerts}


@router.post("/picklists/{no}/pdi-complete")
def pdi_complete(no: str, request: Request, p: Principal = Depends(need("PDI_CHECK")), db: Session = Depends(get_db)):
    k = db.get(models.PickList, no)
    if not k:
        raise HTTPException(404, "Not found")
    p.require_plant(k.plant_code)
    res = _rule(lambda: lpnsvc.pdi_complete(db, k, p.user_id))
    audit(db, p, "PDI_COMPLETE", "picklist", no, None, res, request)
    db.commit()
    return row(k) | res


# ---------------------------------------------------------------- EBS sub-inventory transactions
@router.get("/lpn-txns")
def list_txns(plant: str | None = None, status: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(models.LpnTxn).order_by(models.LpnTxn.created_at.desc())
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter_by(plant_code=pl)
    if status:
        q = q.filter(models.LpnTxn.status.in_(status.split(",")))
    return [row(t) for t in q.limit(1000).all()]


@router.post("/lpn-txns/retry")
def retry_txns(plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    res = lpnsvc.process_pending_txns(db)
    db.commit()
    return res


# ---------------------------------------------------------------- in-plant movement / receipt (web; HHT uses sync events)
class MoveIn(BaseModel):
    scanned: list[str]
    to_zone: str
    load: str | None = None
    remarks: str | None = None


def _pallet_of(db, sc):
    _, pal = services.resolve_tag(db, sc)
    if not pal:
        l = db.get(models.Lpn, sc.strip())
        pal = db.get(models.Pallet, l.pallet_no) if l and l.pallet_no else None
    return pal


@router.post("/moves")
def move_pallets(body: MoveIn, request: Request, p: Principal = Depends(need("INTERNAL_MOVE")), db: Session = Depends(get_db)):
    if not p.plant:
        raise HTTPException(400, "Plant user required")
    out = []
    for sc in body.scanned:
        pal = _pallet_of(db, sc)
        if not pal:
            out.append({"scanned": sc, "ok": False, "message": "Unknown pallet / LPN"}); continue
        sp = db.begin_nested()
        try:
            msg, alerts = lpnsvc.internal_move(db, pal, body.to_zone, p.plant, load=body.load, user_id=p.user_id, remarks=body.remarks)
            sp.commit(); out.append({"scanned": sc, "ok": True, "message": msg, "alerts": alerts})
        except services.RuleError as ex:
            sp.rollback(); out.append({"scanned": sc, "ok": False, "message": str(ex)})
    audit(db, p, "INTERNAL_MOVE", "pallet", body.to_zone, None, {"n": len(body.scanned)}, request)
    db.commit()
    return out


class ReceiptIn(BaseModel):
    scanned: list[str]
    zone: str = "FGWH"
    ref: str | None = None       # sending plant's challan / invoice


@router.post("/plant-receipts")
def plant_receipt(body: ReceiptIn, request: Request, p: Principal = Depends(need("PLANT_RECEIPT")), db: Session = Depends(get_db)):
    if not p.plant:
        raise HTTPException(400, "Plant user required")
    out = []
    for sc in body.scanned:
        _, pal = services.resolve_tag(db, sc)
        if not pal:
            out.append({"scanned": sc, "ok": False, "message": "Unknown pallet"}); continue
        sp = db.begin_nested()
        try:
            msg, alerts = lpnsvc.receive_from_plant(db, pal, body.zone, p.plant, user_id=p.user_id, ref=body.ref)
            sp.commit(); out.append({"scanned": sc, "ok": True, "message": msg, "alerts": alerts})
        except services.RuleError as ex:
            sp.rollback(); out.append({"scanned": sc, "ok": False, "message": str(ex)})
    audit(db, p, "PLANT_RECEIPT", "pallet", body.zone, None, body.model_dump(), request)
    db.commit()
    return out


@router.get("/zones")
def zone_view(plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    """Pallets physically in the plant by zone x loaded/empty, own vs other-plant."""
    pl = _plant(p, plant)
    q = db.query(models.Pallet.zone, models.Pallet.load_state, (models.Pallet.home_plant == pl).label("own"), func.count()).filter(
        models.Pallet.location_plant == pl, models.Pallet.status.notin_(["DISCONTINUED"])).group_by(
        models.Pallet.zone, models.Pallet.load_state, "own")
    grid = {z: {"LOADED": 0, "EMPTY": 0, "other_plant": 0} for z in lpnsvc.ZONES}
    for z, ls, own, n in q.all():
        g = grid.setdefault(z or "YARD", {"LOADED": 0, "EMPTY": 0, "other_plant": 0})
        g[ls or "EMPTY"] += n
        if not own:
            g["other_plant"] += n
    yard_loaded = db.query(models.Pallet).filter(models.Pallet.location_plant == pl, models.Pallet.zone == "YARD",
                                                 models.Pallet.load_state == "LOADED").count()
    return {"plant": pl, "zones": grid, "yard_loaded_violations": yard_loaded, "routes": sorted(f"{a}>{b}" for a, b in lpnsvc.MOVES)}


# ---------------------------------------------------------------- missed scans
@router.get("/scan-misses")
def scan_misses(plant: str | None = None, open_only: bool = True, point: str | None = None,
                p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(models.ScanMiss).order_by(models.ScanMiss.ts.desc())
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter_by(plant_code=pl)
    if open_only:
        q = q.filter_by(resolved=False)
    if point:
        q = q.filter_by(missed_point=point)
    return [row(x) for x in q.limit(1000).all()]


@router.get("/scan-misses/summary")
def scan_miss_summary(plant: str | None = None, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(models.ScanMiss.missed_point, models.ScanMiss.responsible_role, func.count()).filter_by(resolved=False)
    pl = p.scope_plant(plant)
    if pl:
        q = q.filter(models.ScanMiss.plant_code == pl)
    return [{"missed_point": a, "role": b, "open": n} for a, b, n in q.group_by(models.ScanMiss.missed_point, models.ScanMiss.responsible_role).all()]


class ResolveIn(BaseModel):
    resolution: str


@router.post("/scan-misses/{mid}/resolve")
def resolve_miss(mid: int, body: ResolveIn, request: Request, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    m = db.get(models.ScanMiss, mid)
    if not m:
        raise HTTPException(404, "Not found")
    p.require_plant(m.plant_code)
    p.require_any("RECONCILE_CLOSE", "PLANT_USERS_ADMIN", "PICKLIST_CREATE")
    m.resolved, m.resolved_by, m.resolved_at, m.resolution = True, p.user_id, utcnow(), body.resolution
    audit(db, p, "SCAN_MISS_RESOLVE", "scan_miss", str(mid), None, body.model_dump(), request)
    db.commit()
    return row(m)
