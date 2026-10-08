"""Excel master import - reads the Pallet_Tracking_Master_Template sheets."""
from datetime import datetime
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from sqlalchemy.orm import Session
from openpyxl import load_workbook
from .. import models, services
from ..db import get_db
from ..security import need, Principal, audit, hash_pw

router = APIRouter(prefix="/api/v1/import", tags=["import"])


def yn(v) -> bool:
    return str(v).strip().upper() in ("Y", "YES", "TRUE", "1")


def s(v):
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def i(v):
    try:
        return int(float(v)) if v not in (None, "") else None
    except Exception:
        return None


def d(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(v).strip(), fmt)
        except Exception:
            pass
    return None


def rows_of(ws):
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not any(x not in (None, "") for x in r):
            continue
        rec = dict(zip(hdr, r))
        if str(rec.get("Notes", "")).strip().upper() == "EXAMPLE":
            continue
        yield rec


@router.post("/masters")
async def import_masters(request: Request, file: UploadFile = File(...), p: Principal = Depends(need("GLOBAL_MASTERS")), db: Session = Depends(get_db)):
    data = await file.read()
    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except Exception:
        raise HTTPException(400, "Not a valid .xlsx file")
    log: dict[str, dict] = {}
    errors: list[str] = []

    def res(sheet, key):
        log.setdefault(sheet, {"inserted": 0, "updated": 0, "errors": 0})[key] += 1

    # 1 Plants
    if "1_Plants" in wb.sheetnames:
        for r in rows_of(wb["1_Plants"]):
            code = s(r.get("Plant_Code"))
            if not code:
                continue
            code = code.upper()
            obj = db.get(models.Plant, code); new = obj is None
            if new:
                obj = models.Plant(code=code); db.add(obj)
            obj.name = s(r.get("Plant_Name")) or code; obj.city = s(r.get("City")); obj.state = s(r.get("State"))
            obj.gstin = s(r.get("GSTIN")); obj.address = s(r.get("Address")); obj.pallet_prefix = s(r.get("Pallet_No_Prefix"))
            obj.tag_from = i(r.get("Tag_Series_From")); obj.tag_to = i(r.get("Tag_Series_To"))
            obj.challan_source = (s(r.get("Challan_Source")) or "APP").upper(); obj.app_challan_series = s(r.get("App_Challan_Series"))
            obj.accept_other_plant_default = yn(r.get("Accept_Other_Plant_Default")); obj.ebs_org_id = s(r.get("EBS_Org_ID"))
            obj.gcs_endpoint = s(r.get("GCS_Endpoint")); obj.holding_days_default = i(r.get("Holding_Limit_Days_Default")) or 30
            obj.active = yn(r.get("Active")) if r.get("Active") not in (None, "") else True; obj.notes = s(r.get("Notes"))
            res("1_Plants", "inserted" if new else "updated")
        db.flush()
    plants = {x.code for x in db.query(models.Plant).all()}
    roles = {x.code for x in db.query(models.Role).all()}

    # 3 Users
    if "3_Users" in wb.sheetnames:
        for r in rows_of(wb["3_Users"]):
            uid = s(r.get("User_ID"))
            if not uid:
                continue
            uid = uid.lower()
            rl = [x.strip().upper() for x in str(r.get("Roles") or "").replace(",", ";").split(";") if x.strip()]
            bad = [x for x in rl if x not in roles]
            pc = s(r.get("Plant_Code")); pc = pc.upper() if pc else None
            if bad or (pc and pc not in plants) or not rl:
                errors.append(f"3_Users {uid}: bad roles {bad} or plant {pc}"); res("3_Users", "errors"); continue
            u = db.query(models.User).filter_by(user_id=uid).first(); new = u is None
            if new:
                u = models.User(user_id=uid, password_hash=hash_pw("Welcome@123"), must_change_pw=True); db.add(u)
            u.full_name = s(r.get("Full_Name")) or uid; u.plant_code = pc; u.email = s(r.get("Email")); u.mobile = s(r.get("Mobile"))
            u.supervisor_allowed = yn(r.get("Supervisor_PIN_Allowed")); u.device_id = s(r.get("Device_ID"))
            u.active = yn(r.get("Active")) if r.get("Active") not in (None, "") else True; u.notes = s(r.get("Notes"))
            u.roles = [models.UserRole(role_code=x) for x in sorted(set(rl))]
            res("3_Users", "inserted" if new else "updated")

    # 4 Customers
    if "4_Customers" in wb.sheetnames:
        for r in rows_of(wb["4_Customers"]):
            code = s(r.get("Customer_Code")); pc = s(r.get("Plant_Code"))
            if not code or not pc:
                continue
            pc = pc.upper()
            if pc not in plants:
                errors.append(f"4_Customers {code}: plant {pc} unknown"); res("4_Customers", "errors"); continue
            c = db.query(models.Customer).filter_by(code=code, plant_code=pc).first(); new = c is None
            if new:
                c = models.Customer(code=code, plant_code=pc); db.add(c)
            c.name = s(r.get("Customer_Name")) or code; c.ctype = s(r.get("Type")); c.ship_to = s(r.get("Ship_To_Site"))
            c.city = s(r.get("City")); c.state = s(r.get("State")); c.gstin = s(r.get("GSTIN"))
            c.email_to = s(r.get("Email_To")); c.email_cc = s(r.get("Email_CC")); c.contact_name = s(r.get("Contact_Name"))
            c.contact_mobile = s(r.get("Contact_Mobile")); c.holding_limit_days = i(r.get("Holding_Limit_Days"))
            c.reminder_every_days = i(r.get("Reminder_Every_Days")); c.escalation_days = i(r.get("Escalation_Days"))
            c.chargeable = yn(r.get("Chargeable")) if r.get("Chargeable") not in (None, "") else True
            c.return_mode = (s(r.get("Return_Mode")) or "B").upper()[:1]; c.challan_source_override = s(r.get("Challan_Source_Override"))
            c.active = yn(r.get("Active")) if r.get("Active") not in (None, "") else True; c.notes = s(r.get("Notes"))
            res("4_Customers", "inserted" if new else "updated")

    # 5 Transporters
    if "5_Transporters" in wb.sheetnames:
        for r in rows_of(wb["5_Transporters"]):
            code = s(r.get("Transporter_Code"))
            if not code:
                continue
            t = db.get(models.Transporter, code); new = t is None
            if new:
                t = models.Transporter(code=code); db.add(t)
            t.name = s(r.get("Transporter_Name")) or code; t.gstin = s(r.get("GSTIN")); t.contact_name = s(r.get("Contact_Name"))
            t.contact_mobile = s(r.get("Contact_Mobile")); t.email = s(r.get("Email")); t.plant_codes = s(r.get("Plant_Code"))
            t.vehicles = s(r.get("Vehicle_Numbers")); t.active = yn(r.get("Active")) if r.get("Active") not in (None, "") else True
            t.notes = s(r.get("Notes"))
            res("5_Transporters", "inserted" if new else "updated")

    # 7 Tags (before pallets so tags can be assigned)
    if "7_Tags" in wb.sheetnames:
        for r in rows_of(wb["7_Tags"]):
            bn = s(r.get("Batch_No")); pc = s(r.get("Plant_Code")); f = i(r.get("Tag_From")); t = i(r.get("Tag_To"))
            if not bn or not pc or f is None or t is None:
                continue
            pc = pc.upper()
            if pc not in plants or t < f or t - f > 10000:
                errors.append(f"7_Tags {bn}: bad plant/range"); res("7_Tags", "errors"); continue
            b = db.get(models.TagBatch, bn); new = b is None
            if new:
                b = models.TagBatch(batch_no=bn, plant_code=pc, tag_from=f, tag_to=t); db.add(b)
            b.vendor = s(r.get("Vendor")); b.received_date = d(r.get("Received_Date")); b.received_by = s(r.get("Received_By"))
            b.status = (s(r.get("Status")) or "IN_STOCK").upper(); b.notes = s(r.get("Notes"))
            db.flush()
            have = {x[0] for x in db.query(models.Tag.tag_no).filter(models.Tag.batch_no == bn).all()}
            db.bulk_save_objects([models.Tag(tag_no=str(k), plant_code=pc, batch_no=bn) for k in range(f, t + 1) if str(k) not in have])
            res("7_Tags", "inserted" if new else "updated")
        db.flush()

    # 6 Pallets (opening census)
    if "6_Pallets" in wb.sheetnames:
        for r in rows_of(wb["6_Pallets"]):
            no = s(r.get("Pallet_No")); hp = s(r.get("Home_Plant"))
            if not no or not hp:
                continue
            hp = hp.upper()
            if hp not in plants:
                errors.append(f"6_Pallets {no}: plant {hp} unknown"); res("6_Pallets", "errors"); continue
            pal = db.get(models.Pallet, no); new = pal is None
            if new:
                pal = models.Pallet(pallet_no=no, home_plant=hp); db.add(pal)
                db.add(models.PalletHistory(pallet_no=no, event_type="IMPORT", to_status=(s(r.get("Status")) or "AVAILABLE").upper(), plant_code=hp, user_id=p.user_id))
            pal.home_plant = hp; pal.pallet_type = s(r.get("Pallet_Type")); pal.size = s(r.get("Size_LxWxH_mm"))
            try:
                pal.tare_kg = float(r.get("Tare_Weight_kg")) if r.get("Tare_Weight_kg") not in (None, "") else None
            except Exception:
                pal.tare_kg = None
            st = (s(r.get("Status")) or "AVAILABLE").upper(); pal.status = st
            loc = s(r.get("Current_Location"))
            if st == "AT_CUSTOMER":
                pal.customer_code = loc; pal.location_plant = None; pal.dispatch_date = d(r.get("Dispatch_Date")) or datetime.utcnow()
                pal.challan_ref = s(r.get("Challan_Ref"))
            else:
                pal.location_plant = (loc or hp).upper(); pal.customer_code = None
            pal.purchase_date = d(r.get("Purchase_Date")); pal.notes = s(r.get("Notes"))
            db.flush()
            tg = s(r.get("Current_Tag_No"))
            if tg and pal.current_tag != tg:
                tag = db.get(models.Tag, tg)
                if not tag:
                    tag = models.Tag(tag_no=tg, plant_code=hp); db.add(tag)   # tag not in a batch: accept census tag
                if tag.status == "ASSIGNED" and tag.pallet_no != no:
                    errors.append(f"6_Pallets {no}: tag {tg} already on {tag.pallet_no}"); res("6_Pallets", "errors"); continue
                if pal.current_tag:
                    old = db.get(models.Tag, pal.current_tag)
                    if old:
                        old.status = "RETIRED"; old.retired_at = datetime.utcnow()
                tag.status = "ASSIGNED"; tag.pallet_no = no; tag.assigned_at = datetime.utcnow(); tag.plant_code = hp
                db.add(models.TagHistory(pallet_no=no, old_tag=pal.current_tag, new_tag=tg, reason="IMPORT", user_id=p.user_id))
                pal.current_tag = tg
            res("6_Pallets", "inserted" if new else "updated")

    audit(db, p, "IMPORT_MASTERS", "file", file.filename, None, {"summary": log, "errors": errors[:50]}, request)
    db.commit()
    return {"summary": log, "errors": errors}
